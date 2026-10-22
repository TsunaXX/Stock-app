"""Evidence-based simulation on existing quotes; one local writer, atomic cloud merge."""
import hashlib
import io
import json
import math
import queue
import sqlite3
import threading
import time
from contextlib import closing, contextmanager
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP, ROUND_CEILING, ROUND_FLOOR
from zoneinfo import ZoneInfo

import pandas as pd
import requests

VERSION = 'V1.2'
PROTECTION_POLICY = 'balanced-1R-v1'
TZ = ZoneInfo('Asia/Taipei')
ACTIVE = {'等待進場', '模擬持倉', '監控中斷', '資料不足'}
DEFAULT_COSTS = {'stock_shares': 1000, 'stock_discount': 2.8, 'stock_min_fee': 20,
                 'futures_contracts': 1, 'futures_fee': 10, 'slippage_ticks': 1}


def normalize_costs(values):
    values = values if isinstance(values,dict) else {}
    result = {}
    for key,default in DEFAULT_COSTS.items():
        value = number(values.get(key, default))
        minimum = .1 if key == 'stock_discount' else (1 if key in ('stock_shares','futures_contracts') else 0)
        maximum = 10 if key == 'stock_discount' else 10000000
        value = value if value is not None and minimum <= value <= maximum else default
        result[key] = value if key == 'stock_discount' else int(value)
    return result


def stamp(value):
    try:
        dt = pd.Timestamp(value)
        if pd.isna(dt):
            return None
        return dt.tz_localize('Asia/Taipei') if dt.tzinfo is None else dt.tz_convert('Asia/Taipei')
    except (ValueError, TypeError, OverflowError):
        return None


def number(value):
    try:
        n = float(value)
        return n if math.isfinite(n) else None
    except (ValueError, TypeError):
        return None


def option_tick(price):
    """TXO premium ticks, not the underlying index tick."""
    return next((tick for ceiling, tick in ((10, .1), (50, .5), (500, 1), (1000, 5))
                 if price < ceiling), 10)


def option_fills(record, quote, opening):
    prices = quote.get('權利金證據', {})
    fills = {}
    for leg in record.get('選擇權腿', []):
        evidence = prices.get(leg['代碼'], {})
        price, source, trigger_time = number(evidence.get('價格')), stamp(evidence.get('來源時間')), stamp(quote.get('來源時間'))
        if (price is None or price <= 0 or source is None or trigger_time is None
                or not 0 <= (trigger_time - source).total_seconds() <= 2):
            return None
        sign = leg['買賣'] * (1 if opening else -1)
        fill = max(0, price + sign * option_tick(price) * record['成本設定']['slippage_ticks'])
        step = option_tick(fill)
        fills[leg['代碼']] = round((math.ceil(fill / step - 1e-9) if sign > 0 else math.floor(fill / step + 1e-9)) * step, 2)
    if len(fills) == 2:
        credit = -sum(leg['買賣'] * fills[leg['代碼']] for leg in record['選擇權腿'])
        raw_credit = -sum(leg['買賣'] * prices[leg['代碼']]['價格'] for leg in record['選擇權腿'])
        width = number(record.get('限定風險寬度'))
        if width is None or not 0 <= raw_credit <= width or (opening and not 0 <= credit < width):
            return None
    return fills


def trade_id(signal):
    if signal.get('識別版本') == 2 or str(signal.get('策略版本', '')).startswith('V1.2'):
        return hashlib.sha256(('execution-v2|' + plan_id(signal) + '|' + str(signal.get('前次執行ID', ''))).encode()).hexdigest()
    # Decimal spellings and different device receipt times cannot change identity.
    fields = [signal.get(k, '') for k in ('交易日', '市場', '商品鍵', '策略', '策略版本', '方向')]
    if signal.get('選擇權腿'):
        fields.append(json.dumps(signal['選擇權腿'], sort_keys=True, ensure_ascii=False))
    fields += [format(float(signal[k]), '.10g') for k in ('進場價', '停損價', '目標價')]
    return hashlib.sha256('|'.join(map(str, fields)).encode()).hexdigest()


def plan_labels(row):
    source = row.get('計畫來源')
    kind = row.get('計畫類型')
    if not source:
        if row.get('選擇權腿'):
            source = '選擇權操作計畫'
        elif kind in ('日線／15分K', '短波5分K') or str(row.get('商品鍵', '')).endswith(('|日線／15分K', '|短波5分K')):
            source = '指數操作計畫'
        elif row.get('市場') in ('股票', '期貨') and not row.get('觸發商品'):
            source = '股期戰略室'
        else:
            source = '未確認'
    if not kind:
        kind = (str(row.get('市場', '')) + str(row.get('策略', '')) if source == '股期戰略室'
                else str(row.get('商品鍵', '')).rsplit('|', 1)[-1] if source == '指數操作計畫' else '未確認')
    return {'計畫來源': source, '計畫類型': kind}


def plan_id(signal):
    fields = {key: signal.get(key, '') for key in ('市場', '商品鍵', '策略', '策略版本', '方向', '到期時間')}
    fields.update(plan_labels(signal))
    fields['交易日'] = signal.get('交易日', '') if signal.get('策略') == '當沖' else ''
    fields['選擇權腿'] = sorted(signal.get('選擇權腿', []), key=lambda leg: (leg['代碼'], leg['買賣']))
    return hashlib.sha256(('plan-v2|' + json.dumps(fields, ensure_ascii=False, sort_keys=True, separators=(',', ':'))).encode()).hexdigest()


def trade_slot(row):
    fields = [str(row.get(key, '')) for key in ('商品鍵', '市場', '策略', '策略版本', '方向')]
    if row.get('識別版本') == 2:
        fields += list(plan_labels(row).values())
    return '|'.join(fields)


def plan_scope(row):
    return '|'.join(str(row.get(key, '')) for key in ('市場', '商品鍵', '策略', '策略版本')) + '|' + '|'.join(plan_labels(row).values())


def waiting(record, ts, category, reason):
    changed = record.get('最後進場等待原因') != reason or record.get('未進場分類') != category
    record.update({'來源時間': ts, '進場等待原因': reason, '最後進場等待原因': reason,
                   '最後進場等待時間': ts, '未進場分類': category})
    if changed:
        event(record, record['狀態'], ts, reason)


def event_count(row):
    archive = row.get('事件封存', {})
    return archive.get('事件數', 0) - archive.get('保留事件數', 0) + sum(value.get('次數', 1) for value in row.get('事件', []))


def merge_trades(remote, local):
    merged = {}
    for row in list(remote or []) + list(local or []):
        if not isinstance(row, dict) or not row.get('交易ID'):
            continue
        row = dict(row)
        old = merged.get(row['交易ID'])
        if old:
            incoming = row
            if old.get('事件封存') != incoming.get('事件封存') and (old.get('事件封存') or incoming.get('事件封存')):
                raise ValueError('事件封存格式須先核對；保留本機待同步')
            same_stream = incoming.get('串流識別') == old.get('串流識別')
            newer = (str(incoming.get('來源時間', '')), (number(incoming.get('最後序號')) or 0) if same_stream else 0, len(incoming.get('事件', []))) > (
                str(old.get('來源時間', '')), (number(old.get('最後序號')) or 0) if same_stream else 0, len(old.get('事件', [])))
            terminal = incoming.get('狀態') == '已平倉'
            old_terminal = old.get('狀態') == '已平倉'
            if terminal and old_terminal:
                chosen = min((old, incoming), key=lambda r: (bool(r.get('資料缺口')), str(r.get('結案時間', r.get('來源時間',''))), str(r.get('進場時間',''))))
            else:
                chosen = incoming if (newer and not old_terminal) or (terminal and not old_terminal) else old
            row = dict(chosen)
            row['資料缺口'] = bool(row.get('資料缺口')) if row.get('狀態') == '已平倉' else bool(old.get('資料缺口') or incoming.get('資料缺口'))
            if old.get('模擬進場價') is not None and incoming.get('模擬進場價') is not None and old['模擬進場價'] != incoming['模擬進場價']:
                row['資料缺口'] = True
                row['異常原因'] = '跨裝置模擬成交證據不一致，排除有效績效'
            if old.get('成本設定') != incoming.get('成本設定'):
                row['資料缺口'] = True
                row['異常原因'] = '跨裝置成本設定不一致，排除有效績效'
            first = min((old, incoming), key=lambda r: str(r.get('建立時間', '')))
            for key in ('建立時間', '指標快照', '成本設定', '訊號價', '市場環境', '觸發條件', '條件快照'):
                if key in first:
                    row[key] = first[key]
            if row.get('識別版本') == 2:
                for key in ('進場價', '停損價', '目標價'):
                    row[key] = first[key]
            for key in ('最後進場等待原因', '最後進場等待時間', '訊號失效原因', '訊號失效時間', '未進場分類'):
                value = incoming.get(key) or old.get(key)
                if not row.get(key) and value is not None:
                    row[key] = value
            causes = list(dict.fromkeys([value for item in (old, incoming) for value in item.get('資料缺口原因', [])]))
            if causes:
                row['資料缺口原因'] = causes
            events = {}
            for item in (old, incoming):
                for value in item.get('事件', []):
                    if not value.get('事件ID'):
                        continue
                    prior = events.get(value['事件ID'])
                    if prior is None or (str(value.get('結束時間', value['時間'])), value.get('次數', 1)) >= (str(prior.get('結束時間', prior['時間'])), prior.get('次數', 1)):
                        events[value['事件ID']] = value
            row['事件'] = sorted(events.values(), key=lambda e: e['時間'])
            row['事件數'] = event_count(row)
        row.update(plan_labels(row))
        if row.get('進場等待原因') and not row.get('最後進場等待原因'):
            row['最後進場等待原因'] = row['進場等待原因']
        for value in reversed(row.get('事件', [])):
            if value.get('狀態') == '訊號失效' and value.get('原因'):
                row.setdefault('訊號失效原因', value['原因'])
                row.setdefault('訊號失效時間', value['時間'])
                break
        if row.get('識別版本') == 2 and old:
            row['同步確認'] = bool(old.get('同步確認') or incoming.get('同步確認'))
            for key in ('重入解除', '重入來源時間', '重入最後序號', '重入串流識別', '重入等待原因'):
                candidates = [item for item in (old, incoming) if key in item]
                if candidates:
                    row[key] = max(candidates, key=lambda item: str(item.get('重入來源時間', '')))[key]
        restore_profit_protection(row)
        merged[row['交易ID']] = row
    return list(merged.values())


def event(record, state, ts, reason='', evidence=None):
    record['狀態'] = state
    record['來源時間'] = ts
    record['異常原因'] = reason
    if state == '訊號失效':
        record.update({'訊號失效原因': reason, '訊號失效時間': ts})
        if record.get('模擬進場價') is None:
            record['未進場分類'] = '超過有效時間' if '結束' in reason or '到期' in reason else '訊號失效'
    if record.get('資料缺口') and reason and state in ('監控中斷', '資料不足'):
        causes = record.setdefault('資料缺口原因', [])
        if reason not in causes:
            causes.append(reason)
    if state in ('監控中斷', '資料不足') and record.get('模擬進場價') is None:
        record['未進場分類'] = '行情不完整' if record.get('資料缺口') else '暖機等待'
    protection = {key: record[key] for key in ('保利價', '保利淨額', '保利最大浮盈', '初始風險') if key in record} if reason.startswith('保利') else {}
    identity = '|'.join([record['交易ID'], state, ts, reason]) + (json.dumps(protection, sort_keys=True) if protection else '')
    if evidence:
        identity += json.dumps(evidence, ensure_ascii=False, sort_keys=True)
    eid = hashlib.sha256(identity.encode()).hexdigest()
    events = record.setdefault('事件', [])
    diagnostic = record.get('識別版本') == 2 and reason and not protection and state in ('監控中斷', '資料不足', '等待進場')
    last = events[-1] if events else None
    if diagnostic and last and last['狀態'] == state and last.get('原因') == reason and last.get('缺漏證據') == evidence:
        if ts > last.get('結束時間', last['時間']):
            last.update({'起始時間': last['時間'], '結束時間': ts, '次數': last.get('次數', 1) + 1})
        record['事件數'] = event_count(record)
        return
    if not any(e.get('事件ID') == eid for e in record.setdefault('事件', [])):
        item = {'事件ID': eid, '時間': ts, '狀態': state, '原因': reason, **protection}
        if diagnostic:
            item.update({'起始時間': ts, '結束時間': ts, '次數': 1})
        if evidence:
            item['缺漏證據'] = evidence
        events.append(item)
        record['事件數'] = event_count(record)


def new_trade(signal):
    entry,stop,target = (number(signal.get(k)) for k in ('進場價','停損價','目標價'))
    long = signal.get('方向') in ('多頭','偏多')
    if None in (entry,stop,target) or min(entry,stop,target) <= 0 or not (stop < entry < target if long else target < entry < stop):
        raise ValueError('模擬計畫缺少有效進場、停損或目標')
    if signal.get('方向') not in ('多頭','偏多','空頭','偏空') or signal.get('策略') not in ('當沖','波段') or not signal.get('策略版本') or stamp(signal.get('交易日')) is None:
        raise ValueError('訊號方向、策略版本或交易日無效')
    if not signal.get('商品鍵') or signal.get('市場') not in ('股票','期貨') or stamp(signal.get('來源時間')) is None:
        raise ValueError('訊號商品或來源時間無效')
    legs = signal.get('選擇權腿')
    if legs is not None and (not isinstance(legs, list) or len(legs) not in (1, 2)
            or any(not isinstance(leg, dict) or not leg.get('代碼') or leg.get('買賣') not in (-1, 1) for leg in legs)
            or len({leg['代碼'] for leg in legs}) != len(legs)
            or sum(leg['買賣'] for leg in legs) != (1 if len(legs) == 1 else 0)
            or (len(legs) == 2 and number(signal.get('限定風險寬度')) not in (50, 100))
            or signal.get('乘數') != 50 or stamp(signal.get('到期時間')) is None):
        raise ValueError('選擇權腿、乘數或到期時間無效')
    row = {**signal, '交易ID': trade_id(signal), '紀錄類型': '自動', '資料缺口': False,
           'MFE(R)': 0.0, 'MAE(R)': 0.0, '事件': []}
    row.update(plan_labels(row))
    if str(signal.get('策略版本', '')).startswith('V1.2') or signal.get('識別版本') == 2:
        previous = str(signal.get('前次執行ID', ''))
        ordinal = signal.get('進場次數', 1)
        if (previous and (len(previous) != 64 or any(c not in '0123456789abcdef' for c in previous))) or not isinstance(ordinal, int) or ordinal < 1:
            raise ValueError('交易執行鏈或進場次數無效')
        row.update({'原始計畫ID': plan_id(signal), '交易執行ID': row['交易ID'], '前次執行ID': previous,
                    '進場次數': ordinal, '識別版本': 2, '同步確認': False})
    row['成本設定'] = normalize_costs(row.get('成本設定'))
    if row.get('保利規則') == PROTECTION_POLICY:
        row['保利狀態'] = '待進場'
    row['建立時間'] = row['來源時間']
    event(row, '等待進場', row['來源時間'])
    return row


def entry_quality_reason(row, price):
    """Check executable prices and both-side costs; never alter an existing position."""
    if not row.get('進場品質檢查') or row.get('選擇權腿') or row.get('策略') != '當沖':
        return ''
    sign = 1 if row['方向'] in ('多頭', '偏多') else -1
    entry, stop, target = (float(row[k]) for k in ('進場價', '停損價', '目標價'))
    if sign * (price - entry) > abs(entry - stop) * .25:
        return '已偏離原進場點超過 0.25R，等待回測'
    costs = normalize_costs(row.get('成本設定'))
    slip = row['跳動點'] * costs['slippage_ticks']
    fill, exit_target, exit_stop = price + sign * slip, target - sign * slip, stop - sign * slip
    qty = costs['stock_shares'] if row['市場'] == '股票' else costs['futures_contracts']
    multiplier = row['乘數']
    def net(exit_price):
        if row['市場'] == '股票':
            fee = sum(max(costs['stock_min_fee'], math.floor(p * qty * .001425 * costs['stock_discount'] / 10)) for p in (fill, exit_price))
            tax = math.floor((exit_price if sign > 0 else fill) * qty * .0015)
        else:
            rate = number(row.get('期交稅率'))
            if rate is None:
                return None
            fee = costs['futures_fee'] * qty * 2
            tax = sum(round(p * multiplier * qty * rate) for p in (fill, exit_price))
        return sign * (exit_price - fill) * multiplier * qty - fee - tax
    reward, loss = net(exit_target), net(exit_stop)
    if reward is None or loss is None:
        return '未確認交易成本，等待資料'
    if reward <= 0 or loss >= 0 or reward / -loss < 1:
        return '扣除雙邊滑價與成本後風報比不足 1，等待回測'
    return ''


def trade_pnl(row, exit_price, fills=None):
    """One cost definition for simulated exits and the cost-covering protection floor."""
    costs, legs = row['成本設定'], row.get('選擇權腿')
    qty = costs['stock_shares'] if row['市場'] == '股票' else costs['futures_contracts']
    multiplier, entry = row['乘數'], row['模擬進場價']
    sign = 1 if row['方向'] in ('多頭', '偏多') else -1
    gross = sign * (exit_price - entry) * multiplier * qty
    if row['市場'] == '股票':
        fee = sum(max(costs['stock_min_fee'], math.floor(p * qty * .001425 * costs['stock_discount'] / 10)) for p in (entry, exit_price))
        tax = math.floor((exit_price if sign > 0 else entry) * qty * (.0015 if row['策略'] == '當沖' else .003))
    elif legs:
        if fills is None:
            return None
        gross = sum(leg['買賣'] * (fills[leg['代碼']] - row['進場權利金'][leg['代碼']]) for leg in legs) * multiplier * qty
        fee = costs['futures_fee'] * qty * 2 * len(legs)
        tax = sum(int((Decimal(str(p)) * Decimal(str(multiplier)) * qty * Decimal('.001')).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
                  for values in (row['進場權利金'], fills) for p in values.values())
    else:
        rate = number(row.get('期交稅率'))
        if rate is None:
            return None
        fee = costs['futures_fee'] * qty * 2
        tax = sum(round(p * multiplier * qty * rate) for p in (entry, exit_price))
    return {'毛損益': round(gross, 2), '交易成本': round(fee + tax, 2), '模擬損益': round(gross - fee - tax, 2)}


def protection_tick(row, price):
    if row['市場'] == '股票':
        return next((step for ceiling, step in ((10, .01), (50, .05), (100, .1), (500, .5), (1000, 1)) if price < ceiling), 5)
    return row['跳動點']


def restore_profit_protection(row):
    """The existing cloud event union also carries monotonic protection across devices."""
    if row.get('保利規則') != PROTECTION_POLICY or row.get('狀態') not in ACTIVE:
        return
    for item in row.get('事件', []):
        if item.get('時間', '') > row.get('來源時間', ''):
            continue
        for key in ('保利價', '保利淨額', '保利最大浮盈'):
            value = number(item.get(key))
            if value is not None:
                old = number(row.get(key))
                row[key] = value if old is None else (min(old, value) if key == '保利價' and row['方向'] in ('空頭', '偏空') else max(old, value))
        if '初始風險' in item:
            row.setdefault('初始風險', item['初始風險'])
        if item.get('保利價') is not None or item.get('保利淨額') is not None:
            row.setdefault('保利啟動時間', item['時間'])
            row['保利更新時間'] = max(row.get('保利更新時間', ''), item['時間'])
    if row.get('保利價') is not None or row.get('保利淨額') is not None:
        row['保利狀態'] = '保利中'
    if row.get('初始風險', 0) > 0:
        row['保利最大浮盈(R)'] = round(row.get('保利最大浮盈', 0) / row['初始風險'], 4)


def update_profit_protection(row, price, text, pnl=None):
    """Only observed post-entry prices can raise protection; initial risk never changes."""
    if row.get('保利規則') != PROTECTION_POLICY or row.get('資料缺口'):
        return
    legs = row.get('選擇權腿')
    sign = 1 if row['方向'] in ('多頭', '偏多') else -1
    entry, risk = row['模擬進場價'], row['初始風險']
    if legs:
        if pnl is None:
            row['保利狀態'] = '權利金待確認'
            return
        peak = max(row.get('保利最大浮盈', 0), pnl['模擬損益'])
        row['保利最大浮盈'] = peak
        row['保利最大浮盈(R)'] = round(peak / risk, 4)
        if peak < risk:
            return
        candidate = max(0, peak - risk)
        old = row.get('保利淨額')
        row['保利淨額'] = max(old or 0, round(candidate, 2))
        changed = old != row['保利淨額']
    else:
        peak = max(row.get('保利最大浮盈', 0), sign * (price - entry))
        if peak == row.get('保利最大浮盈') and row.get('保利價') is not None:
            return
        row['保利最大浮盈'] = peak
        row['保利最大浮盈(R)'] = round(peak / risk, 4)
        if peak < risk:
            return
        slip = row['跳動點'] * row['成本設定']['slippage_ticks']
        # Solve the exact existing fee/tax definition, including adverse exit slippage.
        lo, hi = 0., peak
        for _ in range(40):
            distance = (lo + hi) / 2
            result = trade_pnl(row, entry + sign * (distance - slip))
            if result is None:
                row['保利狀態'] = '成本待確認'
                return
            if result['模擬損益'] >= 0:
                hi = distance
            else:
                lo = distance
        candidate = entry + sign * max(hi, peak - risk)
        step = Decimal(str(protection_tick(row, candidate)))
        candidate = float((Decimal(str(candidate)) / step).to_integral_value(rounding=ROUND_CEILING if sign > 0 else ROUND_FLOOR) * step)
        # A cost floor beyond the observed market is not an executable protection level.
        if sign * (price - candidate) <= 0:
            row['保利狀態'] = '成本待確認'
            return
        old = row.get('保利價')
        row['保利價'] = candidate if old is None else (max(old, candidate) if sign > 0 else min(old, candidate))
        changed = old != row['保利價']
    row['保利狀態'] = '保利中'
    if changed:
        row.setdefault('保利啟動時間', text)
        row['保利更新時間'] = text
        event(row, '模擬持倉', text, '保利啟動' if old is None else '保利移動')


def advance(record, quote):
    """Only ordered observed prices establish fills. Gaps never count as verified wins."""
    row = json.loads(json.dumps(record, ensure_ascii=False))
    if row['狀態'] not in ACTIVE or row.get('到期未確認'):
        return row
    ts, prior = stamp(quote.get('來源時間')), stamp(row.get('來源時間'))
    price = number(quote.get('價格'))
    if ts is None or (prior is not None and ts < prior):
        return row
    restore_profit_protection(row)
    if row.get('到期時間') and ts >= stamp(row['到期時間']):
        row['到期未確認'] = row.get('模擬進場價') is not None
        row['資料缺口'] = bool(row['到期未確認'] or row.get('資料缺口'))
        event(row, '資料不足' if row['到期未確認'] else '訊號失效', ts.isoformat(),
              '契約已到期；沒有結算或到期前平倉證據，不估算損益')
        return row
    if prior is not None and ts == prior and row['狀態'] != '等待進場' and quote.get('序號') == row.get('最後序號') and quote.get('串流識別') == row.get('串流識別'):
        return row
    text = ts.isoformat()
    stream = quote.get('串流識別')
    old_stream = row.get('串流識別')
    if old_stream is not None and stream is not None and old_stream != stream and (old_stream.split(':')[0] == stream.split(':')[0] or not quote.get('銜接確認')):
        row['資料缺口'] = True
        event(row, '監控中斷', text, '串流重新連線；中斷區間未確認')
    if stream is not None:
        row['串流識別'] = stream
    seq = quote.get('序號')
    old_seq = row.get('串流序號', {}).get(stream) if stream else row.get('最後序號')
    if stream == old_stream and seq is not None and old_seq is not None and seq <= old_seq:
        if quote.get('有效'):
            return row
        seq = None  # Record the invalid source without rewinding the last observed evidence cursor.
    if seq is not None and old_seq is not None and seq > old_seq + 1:
        row['資料缺口'] = True
        event(row, '監控中斷', text, '成交證據佇列缺漏，不能確認停利停損先後',
              {'串流識別': stream, '缺漏起序': old_seq + 1, '缺漏迄序': seq - 1, '目前序號': seq})
    if seq is not None:
        row['最後序號'] = seq
        if stream:
            row.setdefault('串流序號', {})[stream] = seq
    if not quote.get('有效', False) or price is None or price <= 0:
        if row['狀態'] != '資料不足':
            row['資料缺口'] = True
            event(row, '資料不足', text, '行情過期、來源失效或成交證據不足')
        return row
    gap_limit = max(30, float(quote.get('更新秒數', 5)) * 3)
    rest_bridge = row.get('策略') == '波段' and quote.get('休市銜接') and stream == old_stream and seq is not None and old_seq is not None and seq == old_seq + 1
    if prior is not None and (ts - prior).total_seconds() > gap_limit and not rest_bridge:
        row['資料缺口'] = True
        event(row, '監控中斷', text, '兩次可驗證行情之間有缺口；不回填當時進出場')
    if row['狀態'] in ('監控中斷', '資料不足'):
        event(row, '模擬持倉' if row.get('模擬進場價') is not None else '等待進場', text,
              '恢復追蹤；缺口仍保留並排除有效績效' if row.get('資料缺口') else '串流成交證據已就緒')
    if row['策略'] == '當沖' and row['交易日'] != quote.get('交易日'):
        row['資料缺口'] = bool(row.get('模擬進場價') is not None or row.get('資料缺口'))
        event(row, '訊號失效', text, '當沖交易日已結束；缺少收盤成交證據時不估算損益')
        return row
    long = row['方向'] in ('多頭', '偏多')
    sign = 1 if long else -1
    entry, stop, target = (float(row[k]) for k in ('進場價', '停損價', '目標價'))
    tick, multiplier = number(row.get('跳動點')), number(row.get('乘數'))
    if tick is None or tick <= 0 or multiplier is None or multiplier <= 0:
        event(row, '資料不足', text, '缺少有效跳動點或契約乘數，無法驗證成本')
        row['資料缺口'] = True
        return row
    option_legs = row.get('選擇權腿')
    slip = 0 if option_legs else tick * row['成本設定']['slippage_ticks']
    if row['狀態'] == '等待進場':
        # Cross-device continuity compares the price and time of the same observed tick.
        row['最新價'] = price
        invalid = price <= stop if long else price >= stop
        if invalid or (row.get('策略') == '當沖' and str(row.get('交易日')) != str(quote.get('交易日'))):
            event(row, '訊號失效', text, '未進場前已越過停損或交易日已結束')
            return row
        if sign * (price - entry) < 0:
            waiting(row, text, '未觸價', '尚未觸及原進場價')
            return row
        if (not quote.get('進場許可', True) or
                (quote.get('許可方向') and quote['許可方向'] != row['方向']) or
                (quote.get('進場許可時間') and ts < stamp(quote['進場許可時間']))):
            waiting(row, text, '策略未放行', '原策略尚未發布有效同向進場許可')
            return row
        if row.get('進場品質檢查') and (
            quote.get('進場品質許可') is False or
            (quote.get('進場品質時間') and ts < stamp(quote['進場品質時間'])) or
            (quote.get('進場品質方向') and quote['進場品質方向'] != row['方向'])
        ):
            waiting(row, text, '策略未放行', '目前條件未放行同向進場，保留原點位追蹤')
            return row
        fill = price + sign * slip
        if sign * (target - fill) <= 0:
            event(row, '訊號失效', text, '首個成交證據已越過目標；不得回填計畫價成交')
            return row
        quality_reason = entry_quality_reason(row, price)
        if quality_reason:
            # Keep the original levels and wait; do not relabel an old loss or move the target.
            waiting(row, text, '品質檢查不符', quality_reason)
            return row
        if option_legs:
            fills = option_fills(row, quote, True)
            if fills is None:
                row['資料缺口'] = True
                event(row, '資料不足', text, '觸及進場但缺少同時段的各腿權利金成交證據')
                return row
            row['進場權利金'] = fills
            row['觸發進場點位'] = price
            fill = sum(leg['買賣'] * fills[leg['代碼']] for leg in option_legs)
        row['模擬進場價'], row['進場時間'], row['最新價'] = fill, text, price
        if row.get('識別版本') == 2:
            row['進場觸發時間'] = text
            row['成交證據識別'] = {key: quote.get(key) for key in ('來源時間', '序號', '串流識別', '價格', '權利金證據') if key in quote}
            if quote.get('進場許可識別'):
                row['進場許可識別'] = quote['進場許可識別']
        if row.get('保利規則') == PROTECTION_POLICY:
            qty = row['成本設定']['stock_shares'] if row['市場'] == '股票' else row['成本設定']['futures_contracts']
            risk = ((fill if len(option_legs) == 1 else row['限定風險寬度'] + fill) * multiplier * qty
                    if option_legs else abs(fill - stop))
            row.update({'初始風險': risk, '初始風險單位': '元' if option_legs else '點',
                        '保利狀態': '未啟動', '保利最大浮盈': 0., '保利最大浮盈(R)': 0.})
        event(row, '模擬持倉', text)
        return row  # The entry sample cannot also prove a later exit.
    fill = row['模擬進場價']
    trigger_fill = row.get('觸發進場點位', fill)
    risk = abs(trigger_fill - stop)
    r = sign * (price - trigger_fill) / risk if risk else 0
    row['MFE(R)'] = max(row['MFE(R)'], r)
    row['MAE(R)'] = max(row['MAE(R)'], -r)
    row['來源時間'], row['最新價'] = text, price
    stopped = price <= stop if long else price >= stop
    targeted = price >= target if long else price <= target
    # OHLC extrema never establish the order of stop and target.
    high, low = number(quote.get('最高')), number(quote.get('最低'))
    observed_stop = (number(row.get('保利價')) or stop) if row.get('保利規則') == PROTECTION_POLICY and not option_legs else stop
    ambiguous = high is not None and low is not None and high >= max(observed_stop, target) and low <= min(observed_stop, target)
    if ambiguous:
        row['資料缺口'] = True
        event(row, '監控中斷', text, '同一根 K 同時觸及停利停損，先後順序未確認')
        return row
    exit_fills = option_fills(row, quote, False) if option_legs else None
    exit_price = (sum(leg['買賣'] * exit_fills[leg['代碼']] for leg in option_legs)
                  if exit_fills is not None else price - sign * slip)
    pnl = trade_pnl(row, exit_price, exit_fills)
    if row.get('保利規則') == PROTECTION_POLICY and option_legs and exit_fills is None:
        row['資料缺口'] = True
        event(row, '資料不足', text, '保利追蹤缺少同時段的各腿權利金成交證據')
        return row
    protected = False
    if row.get('保利規則') == PROTECTION_POLICY:
        if option_legs:
            floor = number(row.get('保利淨額'))
            protected = floor is not None and pnl is not None and pnl['模擬損益'] <= floor
        else:
            floor = number(row.get('保利價'))
            protected = floor is not None and sign * (price - floor) <= 0
    cutoff = ts.time()
    reached_close = (cutoff.hour == 13 and cutoff.minute >= (30 if row['市場'] == '股票' else 45)) or (row['市場'] == '期貨' and cutoff.hour == 5)
    close_session = quote.get('收盤確認', False) and reached_close and row['策略'] == '當沖'
    if stopped or targeted or protected or close_session:
        if option_legs:
            if exit_fills is None:
                row['資料缺口'] = True
                event(row, '資料不足', text, '觸及出場但缺少同時段的各腿權利金成交證據')
                return row
            row['出場權利金'] = exit_fills
        if pnl is None:
            event(row, '資料不足', text, '此商品未確認期交稅率')
            row['資料缺口'] = True
            return row
        if protected:
            row['保利狀態'] = '保利出場'
        row.update({**pnl, '模擬出場價': exit_price, '結案時間': text,
                    '結果(R)': round(sign * (price - trigger_fill) / risk, 4) if option_legs else round(sign * (exit_price - fill) / risk, 4),
                    '出場原因': '保利' if protected else ('停損' if stopped else ('停利' if targeted else '收盤'))})
        event(row, '已平倉', text)
    elif row.get('保利規則') == PROTECTION_POLICY:
        update_profit_protection(row, price, text, pnl)
    return row


def performance(records, state=None):
    data = pd.DataFrame(records)
    closed = [r for r in records if r.get('狀態') == '已平倉' and not r.get('資料缺口')
              and (r.get('識別版本') != 2 or r.get('同步確認'))
              and number(r.get('模擬損益')) is not None]
    closed.sort(key=lambda r: (r.get('結案時間', ''), r.get('交易ID', '')))
    # Only newly completed trades are folded. Cloud corrections/removals rebuild this filtered view.
    state = {} if state is None else state
    fingerprints = [json.dumps({k:r.get(k) for k in ('交易ID','結案時間','模擬損益','結果(R)','出場原因','市場','策略','策略版本','方向','市場環境','指標快照','條件快照','計畫來源','計畫類型')}, ensure_ascii=False, sort_keys=True) for r in closed]
    previous = state.get('fingerprints', [])
    if fingerprints[:len(previous)] != previous:
        state.clear()
        previous = []
    state.setdefault('wins', 0); state.setdefault('losses', 0)
    state.setdefault('profit', 0.); state.setdefault('loss', 0.)
    state.setdefault('equity', 0.); state.setdefault('peak', 0.); state.setdefault('drawdown', 0.)
    state.setdefault('groups', {})
    state.setdefault('r_total', 0.); state.setdefault('r_count', 0); state.setdefault('protected', 0)
    added = closed[len(previous):]
    for r in added:
        value = float(r['模擬損益'])
        state['wins'] += value > 0; state['losses'] += value < 0
        state['profit'] += max(0, value); state['loss'] += min(0, value)
        state['equity'] += value
        state['peak'] = max(state['peak'], state['equity'])
        state['drawdown'] = max(state['drawdown'], state['peak'] - state['equity'])
        if number(r.get('結果(R)')) is not None:
            state['r_total'] += float(r['結果(R)'])
            state['r_count'] += 1
        state['protected'] += r.get('出場原因') == '保利'
        labels = [{k:r.get(k, '未取得') for k in keys} for keys in (('市場','策略','策略版本'), ('方向','市場環境'))]
        labels.append({'計畫來源': r.get('計畫來源', '股期戰略室'), '計畫類型': r.get('計畫類型', r.get('策略'))})
        labels += [{'指標':k} for k,v in r.get('指標快照', {}).items() if v is not None]
        labels += [{'進場條件':k,'判定':'符合' if v else '未符合'} for k,v in r.get('條件快照', {}).items() if isinstance(v,bool)]
        for label in labels:
            key = json.dumps(label, ensure_ascii=False, sort_keys=True)
            group = state['groups'].setdefault(key, {**label, '有效交易數':0, '獲勝':0, '累積模擬損益':0.})
            group['有效交易數'] += 1; group['獲勝'] += value > 0; group['累積模擬損益'] += value
    state['fingerprints'], state['added'] = fingerprints, len(added)
    count, wins, losses = len(closed), state['wins'], state['losses']
    insufficient = sum(bool(r.get('資料缺口') or r.get('到期未確認') or
                            r.get('狀態') in ('資料不足', '監控中斷')) for r in records)
    metrics = {'全部訊號數': len(records), '資料不足數': insufficient,
               '有效交易數': count, '獲勝': wins, '虧損': losses,
               '勝率(%)': wins / count * 100 if count else None,
               '累積模擬損益': state['equity'], '平均獲利': state['profit'] / wins if wins else None,
               '平均虧損': state['loss'] / losses if losses else None,
               '獲利因子': state['profit'] / abs(state['loss']) if losses else None,
               '每筆期望值': state['equity'] / count if count else None, '最大回撤': state['drawdown'],
               '平均R值': state['r_total'] / state['r_count'] if state['r_count'] else None,
               '保利出場比例(%)': state['protected'] / count * 100 if count else None,
               '訊號有效率(%)': count / len(records) * 100 if records else None,
               '資料完整率(%)': sum(not r.get('資料缺口', False) and r.get('狀態') != '資料不足' for r in records) / len(records) * 100 if records else None}
    groups = [{**g, '勝率(%)':g['獲勝']/g['有效交易數']*100} for g in state['groups'].values()]

    return metrics, pd.DataFrame(groups), data


def plan_performance(records):
    plans = {}
    for row in records:
        if row.get('原始計畫ID'):
            plans.setdefault(row['原始計畫ID'], []).append(row)
    result, first, repeated = [], [], []
    for key, rows in plans.items():
        valid = [row for row in rows if row.get('狀態') == '已平倉' and not row.get('資料缺口')
                 and row.get('同步確認') and number(row.get('模擬損益')) is not None]
        first += [row for row in valid if row['進場次數'] == 1]
        repeated += [row for row in valid if row['進場次數'] > 1]
        known = {row['交易ID'] for row in rows}
        complete = all(not row.get('前次執行ID') or row['前次執行ID'] in known for row in rows)
        latest = max(rows, key=lambda row: (row['進場次數'], row['來源時間']))
        triggered = any(row.get('模擬進場價') is not None for row in rows)
        result.append({'原始計畫ID': key, **plan_labels(rows[0]), '策略版本': rows[0]['策略版本'],
                       '商品': str(rows[0].get('代碼', '')) + ' ' + str(rows[0].get('名稱', '')),
                       '已觸發': triggered, '計畫資料完整': complete,
                       '未進場分類': '' if triggered else latest.get('未進場分類', '等待進場'),
                       '訊號失效原因': '' if triggered else latest.get('訊號失效原因', ''),
                       '有效交易數': len(valid), '計畫累積損益': sum(row['模擬損益'] for row in valid) if complete else None,
                       '未確認執行數': sum(bool(row.get('資料缺口') or not row.get('同步確認') or row.get('狀態') in ACTIVE) for row in rows)})
    count = len(result)
    triggered = sum(row['已觸發'] for row in result)
    metrics = {'原始計畫數': count, '計畫觸發率(%)': triggered / count * 100 if count else None,
               '計畫有效率(%)': sum(row['有效交易數'] > 0 for row in result) / triggered * 100 if triggered and all(row['計畫資料完整'] for row in result) else None,
               '首次進場勝率(%)': sum(row['模擬損益'] > 0 for row in first) / len(first) * 100 if first else None,
               '二次以上勝率(%)': sum(row['模擬損益'] > 0 for row in repeated) / len(repeated) * 100 if repeated else None}
    return metrics, pd.DataFrame(result)


def excel_report(records):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.utils import get_column_letter
    light = [{**{key: value for key, value in row.items() if key != '事件'},
              '事件數': sum(e.get('次數', 1) for e in row.get('事件', []))} for row in records]
    metrics, groups, data = performance(light)
    book = Workbook(write_only=True)
    long_sheet = None
    def write(name, columns, rows):
        nonlocal long_sheet
        sheet, page, count = None, 0, 0
        for values in rows:
            if sheet is None or count >= 1_000_000:
                page += 1
                sheet = book.create_sheet(name if page == 1 else name + str(page))
                sheet.freeze_panes = 'A2'
                sheet.auto_filter.ref = f'A1:{get_column_letter(len(columns))}1000000'
                for i, column in enumerate(columns, 1):
                    sheet.column_dimensions[get_column_letter(i)].width = min(45, max(12, len(str(column)) * 2 + 4))
                sheet.append(columns)
                count = 0
            output = []
            for column, value in zip(columns, values):
                if isinstance(value, (dict, list, tuple)):
                    value = json.dumps(value, ensure_ascii=False, default=str)
                if isinstance(value, str) and len(value) > 16000:
                    if long_sheet is None:
                        long_sheet = book.create_sheet('長文字')
                        long_sheet.append(['交易ID', '事件ID', '欄位', '分段', '總段數', '文字'])
                    row = dict(zip(columns, values))
                    parts = [value[i:i+16000] for i in range(0, len(value), 16000)]
                    for index, part in enumerate(parts, 1):
                        cell = WriteOnlyCell(long_sheet, value=part)
                        cell.data_type = 's'
                        long_sheet.append([row.get('交易ID', ''), row.get('事件ID', ''), column, index, len(parts), cell])
                    value = f'完整文字見長文字：{column}（{len(parts)}段）'
                if not isinstance(value, (str, bool, int, float)) and value is not None:
                    value = None if pd.isna(value) else str(value)
                if isinstance(value, float) and not math.isfinite(value):
                    value = None
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = 's'
                output.append(cell)
            sheet.append(output)
            count += 1
        if sheet is None:
            book.create_sheet(name).append(columns)
    for name, frame in [('總覽', pd.DataFrame([metrics])), ('策略與指標', groups), ('歷史明細', data)]:
        write(name, list(frame.columns), frame.itertuples(index=False, name=None))
    event_columns = list(dict.fromkeys(['交易ID'] + [key for row in records for event in row.get('事件', []) for key in event]))
    write('事件', event_columns, (tuple(row['交易ID'] if key == '交易ID' else event.get(key) for key in event_columns)
                                for row in records for event in row.get('事件', [])))
    if any(row.get('原始計畫ID') for row in records):
        _, plans = plan_performance(light)
        write('原始計畫', list(plans.columns), plans.itertuples(index=False, name=None))
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def database_busy(exc):
    return (isinstance(exc, sqlite3.OperationalError)
            and (getattr(exc, 'sqlite_errorcode', 0) & 255) in (sqlite3.SQLITE_BUSY, sqlite3.SQLITE_LOCKED))


class ModelTracker:
    """One bounded writer per app process; SQLite outbox survives restart, monthly cloud scopes merge."""
    def __init__(self, path, url='', sync_seconds=30, start=True, market_open=None):
        self.path, self.url, self.sync_seconds = str(path), url, sync_seconds
        self.market_open = market_open or (lambda row: True)
        self.queue = queue.Queue(maxsize=512)
        self.status = '尚未同步' if url else '僅本機保存；尚未設定 Google Sheet'
        self.synced_at = None
        self.error = ''
        self.revision = 0
        self.closed = False
        self.stop = threading.Event()
        self.overloaded = False
        self.last_activity = time.monotonic()
        self.last_gap_check = 0
        self.leases = {}
        self.primary_streams = {}
        self.source_quotes = {}
        self.entry_permissions = {}
        self.last_sync = 0
        self.month_versions = {}
        self.sync_positions = {}
        self.hydrated = not bool(url)
        self.history_months = frozenset()
        self.archive_requests = frozenset()
        self.write_lock = threading.RLock()
        with closing(self.db()) as db:
            db.execute('PRAGMA journal_mode=WAL')
        with self.write_db() as db:
            db.execute('CREATE TABLE IF NOT EXISTS trades (id TEXT PRIMARY KEY, month TEXT, slot TEXT, state TEXT, body TEXT, dirty INTEGER)')
            db.execute('CREATE INDEX IF NOT EXISTS trades_month ON trades(month)')
            db.execute('CREATE INDEX IF NOT EXISTS trades_state ON trades(state,slot)')
            db.execute('CREATE INDEX IF NOT EXISTS trades_slot ON trades(slot)')
            db.execute("CREATE INDEX IF NOT EXISTS trades_plan ON trades(json_extract(body,'$.原始計畫ID'))")
            db.execute('CREATE TABLE IF NOT EXISTS model_plans (id TEXT PRIMARY KEY, trade_id TEXT, ordinal INTEGER, scope TEXT)')
            db.execute('CREATE INDEX IF NOT EXISTS model_plans_scope ON model_plans(scope)')
            db.execute('CREATE TABLE IF NOT EXISTS model_event_archives (id TEXT PRIMARY KEY, sha TEXT, body TEXT)')
            for raw in db.execute("SELECT body FROM trades WHERE state IN ('等待進場','模擬持倉')").fetchall():
                row = json.loads(raw[0])
                previous_time = row['來源時間']
                row['資料缺口'] = True
                event(row, '監控中斷', datetime.now(TZ).isoformat(), '服務重新啟動；未確認區間不回填')
                row['來源時間'] = previous_time
                self.store(db, row)
        self.thread = None
        self.sync_thread = None
        if start:
            self.thread = threading.Thread(target=self.run, name='model-tracker', daemon=True)
            self.thread.start()
            if self.url:
                self.sync_thread = threading.Thread(target=self.run_sync, name='model-sync', daemon=True)
                self.sync_thread.start()

    def db(self):
        return sqlite3.connect(self.path, timeout=5)

    @contextmanager
    def write_db(self):
        # Only short local transactions share this lock; HTTP and UI reads never hold it.
        with self.write_lock, closing(self.db()) as db, db:
            db.execute('BEGIN IMMEDIATE')
            yield db

    def submit(self, signals, quotes):
        self.last_activity = time.monotonic()
        for key, quote in quotes.items():
            self.leases[key] = (self.last_activity, max(30, float(quote.get('更新秒數',5)) * 3))
        try:
            self.queue.put_nowait((signals, quotes))
            return True
        except queue.Full:
            self.overloaded = True
            self.error = '記錄佇列已滿；保留資料缺口，不阻塞行情'
            return False

    def records(self, start=None, end=None, details=True):
        with closing(self.db()) as db, db:
            projection = 'body' if details else "json_set(json_remove(body,'$.事件'),'$.事件數',COALESCE(json_extract(body,'$.事件數'),json_array_length(body,'$.事件'),0))"
            rows = db.execute(f'SELECT {projection} FROM trades WHERE month >= ? AND month <= ? ORDER BY month,id',
                              (start or '000000', end or '999999')).fetchall()
        result = [json.loads(r[0]) for r in rows]
        for row in result:
            row.update(plan_labels(row))
        return result

    def details(self, ids):
        ids = list(dict.fromkeys(ids))
        result = []
        with closing(self.db()) as db:
            for offset in range(0, len(ids), 200):
                chunk = ids[offset:offset+200]
                result += [json.loads(row[0]) for row in db.execute(
                    'SELECT body FROM trades WHERE id IN (' + ','.join('?' for _ in chunk) + ')', chunk)]
            for row in result:
                reference = row.get('事件封存')
                if reference:
                    saved = db.execute('SELECT body FROM model_event_archives WHERE id=? AND sha=?', (row['交易ID'], reference['sha'])).fetchone()
                    if saved is None:
                        raise ValueError('完整原始事件尚在載入，請稍後再匯出')
                    combined = {event['事件ID']: event for event in json.loads(saved[0])}
                    combined.update({event['事件ID']: event for event in row.get('事件', [])})
                    row['事件'] = sorted(combined.values(), key=lambda event: event['時間'])
        return result

    def events_ready(self, ids, request=False):
        ids = list(dict.fromkeys(ids))
        missing = set()
        with closing(self.db()) as db:
            for offset in range(0, len(ids), 200):
                chunk = ids[offset:offset+200]
                rows = db.execute("SELECT id,json_extract(body,'$.事件封存') FROM trades WHERE id IN (" + ','.join('?' for _ in chunk) + ')', chunk)
                for key, raw in rows:
                    if not raw:
                        continue
                    reference = json.loads(raw)
                    if not db.execute('SELECT 1 FROM model_event_archives WHERE id=? AND sha=?', (key, reference['sha'])).fetchone():
                        missing.add(reference['scope'])
        if request and missing:
            self.archive_requests = self.archive_requests | missing
            self.last_activity = time.monotonic()
        return not missing

    def load_event_archive(self, scope):
        data = self.remote_get(scope)
        if data.get('model_event_archive') != 1 or not isinstance(data.get('records'), list):
            raise ValueError('原始事件封存格式無效；保留資料')
        for record in data['records']:
            text = record.get('events_json', '')
            if hashlib.sha256(text.encode()).hexdigest() != record.get('sha') or not isinstance(json.loads(text), list):
                raise ValueError('原始事件封存校驗不一致；保留資料')
        for offset in range(0, len(data['records']), 20):
            with self.write_db() as db:
                for record in data['records'][offset:offset+20]:
                    db.execute('INSERT OR REPLACE INTO model_event_archives VALUES (?,?,?)', (record['交易ID'], record['sha'], record['events_json']))
        self.archive_requests = self.archive_requests - {scope}
        self.revision += 1

    def events(self, ids, limit=500):
        ids = list(dict.fromkeys(ids))[:100]
        if not ids:
            return []
        if not self.events_ready(ids, request=True):
            return []
        with closing(self.db()) as db:
            query = ','.join('?' for _ in ids)
            archived = db.execute('SELECT a.id,e.value FROM model_event_archives a,json_each(a.body) e WHERE a.id IN (' + query + ') LIMIT ?', [*ids, limit]).fetchall()
            rows = archived
            if len(rows) < limit:
                known = {json.loads(body)['事件ID'] for _, body in rows}
                live = db.execute("SELECT t.id,e.value FROM trades t,json_each(t.body,'$.事件') e WHERE t.id IN (" + query + ') LIMIT ?', [*ids, limit]).fetchall()
                rows += [(key, body) for key, body in live if json.loads(body)['事件ID'] not in known][:limit-len(rows)]
        return [{**json.loads(body), '交易ID': key} for key, body in rows]

    def request_history(self, start, end):
        self.history_months = frozenset(period.strftime('%Y%m') for period in pd.period_range(start, end, freq='M'))
        self.last_activity = time.monotonic()

    def plan_records(self, ids, end):
        ids = list(dict.fromkeys(ids))
        result = []
        with closing(self.db()) as db:
            for offset in range(0, len(ids), 100):
                chunk = ids[offset:offset+100]
                query = "SELECT json_remove(body,'$.事件') FROM trades WHERE json_extract(body,'$.原始計畫ID') IN (" + ','.join('?' for _ in chunk) + ") AND json_extract(body,'$.交易日')<=?"
                result += [json.loads(row[0]) for row in db.execute(query, [*chunk, str(end)])]
        known = {row['交易ID'] for row in result}
        missing = {row['前次執行月份'] for row in result if row.get('前次執行月份') and row.get('前次執行ID') not in known}
        self.history_months = self.history_months | missing
        return result

    def sync_months(self, catalog, pending):
        if catalog.get('model_history') != 1:
            return sorted(set(pending + catalog.get('model_months', [])))
        cutoff = (datetime.now(TZ) - pd.Timedelta(days=90)).strftime('%Y%m')
        active = {month for month, value in catalog.get('model_active_months', {}).items() if value}
        return sorted(set(pending) | (set(self.history_months) & set(catalog.get('model_months', []))) | active |
                      {month for month in catalog.get('model_months', []) if month >= cutoff})

    def active_plans(self, source, day):
        month = day[:7].replace('-', '')
        with closing(self.db()) as db, db:
            rows = db.execute("SELECT body FROM trades WHERE month=? AND state IN ('等待進場','模擬持倉','監控中斷','資料不足') AND json_extract(body,'$.計畫來源')=? AND json_extract(body,'$.交易日')=?", (month, source, day)).fetchall()
        return [row for raw in rows if (row := json.loads(raw[0])).get('計畫來源') == source and row.get('交易日') == day]

    def protection_plans(self, market, keys, mode, day, source=None):
        """Bounded current-plan reads by indexed instrument, without loading performance history."""
        result = {}
        with closing(self.db()) as db:
            for key in keys:
                prefix = f'{key}|{market}|{mode}|'
                raw = db.execute("""SELECT body FROM trades WHERE slot>=? AND slot<?
                    AND json_extract(body, '$.保利規則')=?
                    AND (json_extract(body, '$.計畫來源')=? OR (?='股期戰略室' AND json_extract(body, '$.計畫來源') IS NULL))
                    AND ((?='波段' AND state IN ('等待進場','模擬持倉','監控中斷','資料不足')) OR json_extract(body, '$.交易日')=?)
                    ORDER BY (state IN ('模擬持倉','監控中斷','資料不足') AND json_extract(body, '$.模擬進場價') IS NOT NULL) DESC,
                    rowid DESC LIMIT 1""", (prefix, prefix + '\uffff', PROTECTION_POLICY, source or '股期戰略室', source or '股期戰略室', mode, day)).fetchone()
                if raw:
                    result[str(key)] = json.loads(raw[0])
        return result

    def store(self, db, row, dirty=1):
        row.update(plan_labels(row))
        db.execute('INSERT INTO trades VALUES (?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET state=excluded.state,body=excluded.body,dirty=MAX(trades.dirty,excluded.dirty)',
                   (row['交易ID'], row['交易日'][:7].replace('-', '').replace('/', ''),
                    trade_slot(row),
                    row['狀態'], json.dumps(row, ensure_ascii=False), dirty))
        if row.get('識別版本') == 2:
            db.execute('''INSERT INTO model_plans VALUES (?,?,?,?) ON CONFLICT(id) DO UPDATE
                SET trade_id=excluded.trade_id,ordinal=excluded.ordinal,scope=excluded.scope
                WHERE excluded.ordinal>model_plans.ordinal OR excluded.trade_id=model_plans.trade_id''',
                (row['原始計畫ID'], row['交易ID'], row['進場次數'], plan_scope(row)))

    def process(self, signals, quotes):
        selected = {}
        now = time.monotonic()
        for key, quote in quotes.items():
            stream = quote.get('串流識別')
            if not stream:
                selected[key] = quote
                continue
            sources = self.source_quotes.setdefault(key, {})
            sources[stream] = (now, quote)
            for old in list(sources):
                if now - sources[old][0] > 120:
                    sources.pop(old)
            primary = self.primary_streams.get(key, stream)
            if primary.split(':')[0] == stream.split(':')[0] and primary != stream:
                primary = stream  # Same connection has a new transport/buffer epoch.
            chosen = sources.get(primary)
            if chosen is None or not chosen[1].get('有效') or now - chosen[0] > 30:
                candidates = [(name, value) for name,value in sources.items() if value[1].get('有效') and now-value[0] <= 30]
                if candidates:
                    primary, chosen = max(candidates, key=lambda item: str(item[1][1].get('來源時間','')))
            self.primary_streams[key] = primary
            selected[key] = chosen[1] if chosen else quote
        quotes = selected
        with self.write_db() as db:
            if self.overloaded:
                for raw in db.execute("SELECT body FROM trades WHERE state IN ('等待進場','模擬持倉','監控中斷','資料不足')").fetchall():
                    row = json.loads(raw[0])
                    row['資料缺口'] = True
                    event(row, '監控中斷', datetime.now(TZ).isoformat(), '記錄佇列滿載，成交證據未確認')
                    self.store(db, row)
                    self.revision += 1
                self.overloaded = False
            for signal in signals:
                row = new_trade(signal)
                if row.get('識別版本') == 2 and not self.url:
                    row['同步確認'] = True
                existing = db.execute('SELECT body FROM trades WHERE id=?', (row['交易ID'],)).fetchone()
                # One active simulation per instrument/mode/direction/version; a changing plan is not a new fill.
                slot = trade_slot(row)
                open_rows = db.execute("SELECT body FROM trades WHERE slot=? AND state IN ('等待進場','模擬持倉','監控中斷','資料不足')", (slot,)).fetchall()
                same_open = any(r.get('狀態') in ACTIVE and (row['策略'] != '當沖' or r.get('交易日') == row['交易日']) and all(r.get(k) == row.get(k) for k in ('商品鍵', '市場', '策略', '策略版本', '方向'))
                                for r in (json.loads(x[0]) for x in open_rows))
                if not existing and not same_open:
                    self.store(db, row)
                    self.revision += 1
            active_rows = []
            for key in quotes:
                market, product = key.split('|', 1)
                prefix = product + '|' + market + '|'
                active_rows.extend(db.execute("SELECT body FROM trades WHERE slot>=? AND slot<? AND state IN ('等待進場','模擬持倉','監控中斷','資料不足')", (prefix, prefix + '\uffff')).fetchall())
            for raw in active_rows:
                row = json.loads(raw[0])
                quote = quotes.get(row['市場'] + '|' + row['商品鍵'])
                if quote is None or (quote.get('休市') and not quote.get('收盤確認')) or row['狀態'] not in ACTIVE:
                    continue
                if quote.get('策略') and quote['策略'] != row['策略'] and row.get('模擬進場價') is None:
                    continue
                if (quote.get('暖機等待') and row.get('最後序號') is None and row.get('模擬進場價') is None
                        and (not row.get('到期時間') or stamp(quote['來源時間']) < stamp(row['到期時間']))):
                    if row['狀態'] != '資料不足':
                        event(row, '資料不足', row['來源時間'], '等待第一組串流成交證據，尚未確認進場')
                        self.store(db, row)
                        self.revision += 1
                    continue
                updated = row
                evidence = quote.get('成交證據', [])
                quote = dict(quote)
                prior_time, prior_price = row.get('來源時間'), row.get('最新價', row.get('訊號價'))
                # Stream timestamps normally match verbatim; parse only cross-format bridges.
                quote['銜接確認'] = any(t.get('來源時間') == prior_time and t.get('價格') == prior_price for t in evidence)
                if not quote['銜接確認']:
                    prior_stamp = stamp(prior_time)
                    quote['銜接確認'] = any(t.get('價格') == prior_price and stamp(t.get('來源時間')) == prior_stamp for t in evidence)
                permission = self.entry_permissions.get((row['市場'] + '|' + row['商品鍵'], row['策略']), {})
                permission_time = stamp(permission.get('進場品質時間'))
                current_time = stamp(quote.get('進場品質時間'))
                continuous = (permission.get('有效') and permission.get('交易日') == quote.get('交易日')
                              and permission.get('串流識別') == quote.get('串流識別')
                              and permission_time is not None and current_time is not None
                              and 0 <= (current_time - permission_time).total_seconds() <= max(30, float(quote.get('更新秒數', 5)) * 3))
                for tick in evidence:
                    if str(tick.get('來源時間', '')) < str(updated['來源時間']):
                        continue
                    sample = {**quote, **tick}
                    # A new decision applies from its publication time, never to earlier observed fills.
                    tick_time = stamp(tick.get('來源時間')) if continuous and updated.get('模擬進場價') is None else None
                    if tick_time is not None and permission_time <= tick_time < current_time:
                        sample.update({field: permission[field] for field in ('進場品質許可', '進場品質方向', '進場品質時間')})
                    updated = advance(updated, sample)
                if not evidence or not quote.get('有效'):
                    updated = advance(updated, quote)
                if updated != row:
                    self.store(db, updated)
                    self.revision += 1
            self.reenter(db, quotes)
        for key, quote in quotes.items():
            if quote.get('策略') and quote.get('進場品質時間'):
                identity = (key, quote['策略'])
                previous_time = stamp(self.entry_permissions.get(identity, {}).get('進場品質時間'))
                current_time = stamp(quote['進場品質時間'])
                if current_time is not None and (previous_time is None or current_time >= previous_time):
                    self.entry_permissions[identity] = {
                        field: quote.get(field) for field in ('進場品質許可', '進場品質方向', '進場品質時間', '有效', '交易日', '串流識別')}

    def reenter(self, db, quotes):
        seed_fields = ('交易日', '市場', '商品鍵', '代碼', '名稱', '交易時段', '策略', '策略版本', '方向',
                       '進場價', '停損價', '目標價', '觸發條件', '訊號狀態', '指標快照', '條件快照',
                       '跳動點', '乘數', '期交稅率', '成本設定', '進場品質檢查', '保利規則', '市場環境',
                       '計畫來源', '計畫類型', '觸發商品', '商品類型', '選擇權腿', '到期時間', '限定風險寬度')
        for key, quote in quotes.items():
            if not str(quote.get('策略版本', '')).startswith('V1.2'):
                continue
            market, product = key.split('|', 1)
            scope = plan_scope({**quote, '市場': market, '商品鍵': product})
            plans = db.execute("SELECT t.body FROM model_plans p JOIN trades t ON t.id=p.trade_id WHERE p.scope=? AND t.state='已平倉'", (scope,)).fetchall()
            permission = self.entry_permissions.get((key, quote.get('策略')), {})
            before, current = stamp(permission.get('進場品質時間')), stamp(quote.get('進場品質時間'))
            permit_bridge = (permission.get('有效') and before is not None and current is not None
                             and permission.get('串流識別') == quote.get('串流識別')
                             and 0 <= (current-before).total_seconds() <= max(30, float(quote.get('更新秒數', 5)) * 3))
            for raw, in plans:
                row = json.loads(raw)
                if row['狀態'] != '已平倉' or row.get('資料缺口') or (row['策略'] == '當沖' and row['交易日'] != quote.get('交易日')):
                    continue
                if quote.get('休市') or not quote.get('有效'):
                    if row.get('重入解除'):
                        row['重入解除'] = None
                        row['重入等待原因'] = '行情或時段不足；等待新鮮條件重新解除及成立'
                        self.store(db, row)
                        self.revision += 1
                    continue
                original = json.loads(raw)
                closed_time = stamp(row['結案時間'])
                cursor_time = stamp(row.get('重入來源時間') or row['結案時間'])
                cursor_seq = row.get('重入最後序號', row.get('最後序號'))
                cursor_stream = row.get('重入串流識別', row.get('串流識別'))
                cleared = row.get('重入解除')
                for tick in quote.get('成交證據', []):
                    sample = {**quote, **tick}
                    ts = stamp(sample.get('來源時間'))
                    seq, stream = sample.get('序號'), sample.get('串流識別')
                    if ts is None or ts <= closed_time or (cursor_time is not None and ts < cursor_time):
                        continue
                    if stream == cursor_stream and seq is not None and cursor_seq is not None and seq <= cursor_seq:
                        continue
                    if permit_bridge and before <= ts < current:
                        sample.update({field: permission[field] for field in ('進場品質許可', '進場品質方向', '進場品質時間')})
                    continuous = (cursor_time is not None and 0 <= (ts - cursor_time).total_seconds() <= max(30, float(quote.get('更新秒數', 5)) * 3)
                                  and stream == cursor_stream and seq is not None and cursor_seq is not None and seq == cursor_seq + 1)
                    if not continuous or not sample.get('有效'):
                        cleared = None
                    cursor_time, cursor_seq, cursor_stream = ts, seq, stream
                    if row['狀態'] != '已平倉':
                        row = advance(row, sample)
                        if row['狀態'] == '已平倉':
                            self.store(db, row)
                            closed_time, cleared = stamp(row['結案時間']), None
                        continue
                    row.update({'重入來源時間': ts.isoformat(), '重入最後序號': seq, '重入串流識別': stream})
                    if row.get('到期時間') and ts >= stamp(row['到期時間']):
                        row['重入解除'] = None
                        row['重入等待原因'] = '原契約已到期，不能重新進場'
                        break
                    price = number(sample.get('價格'))
                    if price is None or not sample.get('有效'):
                        row['重入解除'] = None
                        continue
                    direction = sample.get('進場品質方向') or sample.get('許可方向')
                    allowed = sample.get('進場許可', True) and sample.get('進場品質許可', True) and (not direction or direction == row['方向'])
                    hit = price >= row['進場價'] if row['方向'] in ('多頭', '偏多') else price <= row['進場價']
                    permit_time = stamp(sample.get('進場許可時間'))
                    if (sample.get('進場許可識別') and sample['進場許可識別'] != row.get('進場許可識別')
                            and permit_time is not None and closed_time < permit_time <= ts):
                        cleared = {'來源時間': permit_time.isoformat(), '許可識別': sample['進場許可識別'], '型態': '新許可'}
                    if not allowed or not hit:
                        if cleared is None:
                            cleared = {'來源時間': ts.isoformat(), '序號': seq, '串流識別': stream, '價格': price}
                        row['重入等待原因'] = '原進場條件已解除，等待重新成立'
                    elif cleared and continuous and (ts > stamp(cleared['來源時間']) or cleared.get('型態') == '新許可'):
                        seed = {field: row[field] for field in seed_fields if field in row}
                        seed.update({'來源時間': ts.isoformat(), '訊號價': price, '前次執行ID': row['交易ID'],
                                     '前次執行月份': row['交易日'][:7].replace('-', ''),
                                     '進場次數': row['進場次數'] + 1, '交易日': quote['交易日'],
                                     '成本設定': quote.get('成本設定', row['成本設定']), '重入解除證據': cleared})
                        for field in ('指標快照', '條件快照', '市場環境', '觸發條件'):
                            if field in quote:
                                seed[field] = quote[field]
                        candidate = advance(new_trade(seed), sample)
                        if candidate.get('模擬進場價') is not None and not candidate.get('資料缺口'):
                            candidate['同步確認'] = not bool(self.url)
                            row['重入解除'] = cleared
                            self.store(db, row)
                            self.store(db, candidate)
                            self.revision += 1
                            row, cleared = candidate, None
                            continue
                        row['重入等待原因'] = candidate.get('最後進場等待原因') or candidate.get('異常原因') or '等待完整新成交證據'
                    else:
                        row['重入等待原因'] = '等待原進場條件先解除，不因持續觸價重複建倉'
                    row['重入解除'] = cleared
                if row != original:
                    self.store(db, row)
                    self.revision += 1

    def mark_gaps(self):
        if not self.queue.empty():
            return  # Drain already observed evidence before declaring the producer inactive.
        now = time.monotonic()
        with self.write_db() as db:
            for raw in db.execute("SELECT body FROM trades WHERE state IN ('等待進場','模擬持倉')").fetchall():
                row = json.loads(raw[0])
                if row['策略'] == '波段' and not self.market_open(row):
                    continue
                lease = self.leases.get(row['市場'] + '|' + row['商品鍵'])
                if lease is not None and now - lease[0] <= lease[1]:
                    continue
                row['資料缺口'] = True
                event(row, '監控中斷', datetime.now(TZ).isoformat(), '監控停止、登入中斷或頁面離開；未確認區間不回填')
                self.store(db,row)
                self.revision += 1

    def remote_get(self, scope, trade_ids=None):
        for attempt in range(3):
            try:
                params = {'scope': scope, '_ts': time.time_ns()}
                if trade_ids is not None:
                    params['trade_ids'] = ','.join(trade_ids)
                response = requests.get(self.url, params=params,
                                        headers={'Cache-Control': 'no-cache, no-store, max-age=0', 'Pragma': 'no-cache'},
                                        timeout=(5, 35))
                response.raise_for_status()
                break
            except (requests.ConnectionError, requests.Timeout, requests.HTTPError) as exc:
                status = exc.response.status_code if exc.response is not None else None
                if status is not None and status not in (429, 500, 502, 503, 504):
                    raise
                if attempt == 2:
                    reason = f'HTTP {status}' if status else type(exc).__name__
                    raise RuntimeError(f'Google Sheet 暫時無法讀取（{scope}／{reason}，已嘗試 3 次）；保留待同步並自動重試') from exc
                time.sleep(2 ** attempt)
        payload = response.json()
        if payload.get('success') is not True:
            raise ValueError(payload.get('error', 'Google Sheet 讀取失敗'))
        data = payload.get('data') or {}
        if not isinstance(data, dict):
            raise ValueError('Google Sheet 儲存格式不符，保留待同步')
        return data

    def remote_save(self, scope, data, delta=False):
        response = requests.post(self.url, json={'scope': scope, 'data': json.dumps(data, ensure_ascii=False),
                                                'updated_at': datetime.now(TZ).isoformat()}, timeout=(5, 35))
        response.raise_for_status()
        payload = response.json()
        if payload.get('success') is not True:
            raise ValueError(payload.get('error', 'Google Sheet 寫入失敗'))
        actual = self.remote_get(scope, [r['交易ID'] for r in data['model_trades']]) if delta else self.remote_get(scope)
        if actual.get('model_schema') != 1:
            raise ValueError('Apps Script 尚未部署策略追蹤原子合併 V1.0；資料保留待同步')
        if delta:
            version = str(payload.get('model_version', ''))
            previous = str(payload.get('model_previous_version', ''))
            read_version = str(actual.get('model_version', ''))
            if not version.isdigit() or not read_version.isdigit() or int(read_version) < int(version) or (previous and not previous.isdigit()):
                raise ValueError('交易同步版本回讀無效；保留待同步')
            actual['_committed_version'], actual['_previous_version'] = version, previous
        return actual

    def sync(self):
        if not self.url:
            return
        catalog = self.remote_get('strategy_signals')
        if catalog.get('model_schema') != 1:
            raise ValueError('Apps Script 尚未部署策略追蹤原子合併 V1.0；資料保留待同步')
        with closing(self.db()) as db, db:
            pending = [r[0] for r in db.execute('SELECT DISTINCT month FROM trades WHERE dirty=1')]
        versions = catalog.get('model_versions', {})
        if not isinstance(versions, dict) or not isinstance(catalog.get('model_months', []), list) or any(not isinstance(m,str) or len(m) != 6 or not m.isdigit() or not 1 <= int(m[4:]) <= 12 for m in catalog.get('model_months', [])):
            raise ValueError('Google Sheet 月份索引無效；保留待同步')
        if catalog.get('model_history') == 1 and not isinstance(catalog.get('model_active_months'), dict):
            raise ValueError('Google Sheet 活動月份索引無效；保留待同步')
        if catalog.get('model_sync') == 2:
            return self.sync_delta(catalog, pending)
        if catalog.get('model_reentry') != 1:
            with closing(self.db()) as db:
                if db.execute("SELECT 1 FROM trades WHERE dirty=1 AND json_extract(body,'$.識別版本')=2 LIMIT 1").fetchone():
                    raise ValueError('Apps Script 尚未支援 V1.2 執行鏈；新紀錄保留本機待同步')
        months = sorted(set(pending + [m for m in catalog.get('model_months', [])
                                      if not self.hydrated or versions.get(m) != self.month_versions.get(m)]))
        for month in months:
            scope = 'strategy_signals:' + month
            remote = self.remote_get(scope).get('model_trades', [])
            if not isinstance(remote,list):
                raise ValueError('Google Sheet 交易紀錄格式無效')
            for row in remote:
                if not isinstance(row,dict) or new_trade(row)['交易ID'] != row.get('交易ID') or row.get('狀態') not in ACTIVE | {'已平倉','訊號失效'}:
                    raise ValueError('Google Sheet 交易ID或狀態無效；保留待同步')
            local = self.records(month, month)
            local_versions = {r['交易ID']: r for r in local}
            merged = merge_trades(remote, local)
            if month in pending:
                actual = self.remote_save(scope, {'model_schema': 1, 'model_trades': merged})
                saved_rows = actual.get('model_trades', [])
                # Apps Script collation can order identical ISO event times differently.
                saved = {r['交易ID']: r for r in merge_trades(saved_rows, saved_rows)}
                if any(r['交易ID'] not in saved or merge_trades([saved[r['交易ID']]], [r])[0] != saved[r['交易ID']] for r in merged):
                    raise ValueError('交易事件回讀不一致；保留待同步並重試')
                catalog_check = self.remote_get('strategy_signals')
                if month not in catalog_check.get('model_months', []):
                    raise ValueError('月份索引回讀不一致；保留待同步')
                merged = list(saved.values())
            self.merge_synced_rows(merged, local_versions if month in pending else {})
            self.month_versions[month] = versions.get(month)
        self.hydrated = True
        self.synced_at = datetime.now(TZ).strftime('%Y/%m/%d %H:%M:%S')
        self.status, self.error = '已回讀驗證', ''

    def merge_synced_rows(self, rows, sent=None):
        sent = sent or {}
        for row in rows:
            if not isinstance(row, dict) or new_trade(row)['交易ID'] != row.get('交易ID') or row.get('狀態') not in ACTIVE | {'已平倉','訊號失效'}:
                raise ValueError('Google Sheet 交易ID或狀態無效；保留待同步')
        # Migration can meet old local events still waiting to upload; read cold data outside the write lock.
        required = set()
        with closing(self.db()) as db:
            for row in rows:
                reference = row.get('事件封存')
                if reference and db.execute("SELECT 1 FROM trades WHERE id=? AND json_extract(body,'$.事件封存') IS NULL", (row['交易ID'],)).fetchone():
                    if not db.execute('SELECT 1 FROM model_event_archives WHERE id=? AND sha=?', (row['交易ID'], reference['sha'])).fetchone():
                        required.add(reference['scope'])
        for scope in sorted(required):
            self.load_event_archive(scope)
        for offset in range(0, len(rows), 20):
            with self.write_db() as db:
                for row in rows[offset:offset + 20]:
                    raw = db.execute('SELECT body FROM trades WHERE id=?', (row['交易ID'],)).fetchone()
                    current = json.loads(raw[0]) if raw else None
                    if current and row.get('事件封存') and not current.get('事件封存'):
                        reference = row['事件封存']
                        cached = db.execute('SELECT body FROM model_event_archives WHERE id=? AND sha=?', (row['交易ID'], reference['sha'])).fetchone()
                        if cached is None:
                            self.archive_requests = self.archive_requests | {reference['scope']}
                            raise ValueError('封存核對待重試；保留本機待同步')
                        original = {e['事件ID']: e for e in json.loads(cached[0])}
                        hot = {e['事件ID']: e for e in row['事件']}
                        for value in current.get('事件', []):
                            if value['事件ID'] not in original:
                                hot[value['事件ID']] = value
                            elif original[value['事件ID']] != value:
                                raise ValueError('舊事件與封存證據不一致；保留本機待同步')
                        current.update({'事件封存': reference, '事件格式': 2, '事件': list(hot.values())})
                    merged = merge_trades([row], [current] if current else [])[0]
                    missing_labels = any(row.get(key) != merged.get(key) for key in (
                        '計畫來源', '計畫類型', '最後進場等待原因', '訊號失效原因', '訊號失效時間', '事件數', '事件'))
                    self.store(db, merged, dirty=int(missing_labels))
                    if current == sent.get(row['交易ID']) and merged == row:
                        db.execute('UPDATE trades SET dirty=0 WHERE id=?', (row['交易ID'],))
            self.revision += 1

    def sync_delta(self, catalog, pending):
        versions = catalog.get('model_versions', {})
        blocked = False
        for month in self.sync_months(catalog, pending):
            scope = 'strategy_signals:' + month
            if not self.hydrated or versions.get(month) != self.month_versions.get(month):
                remote = self.remote_get(scope)
                if not isinstance(remote.get('model_trades', []), list):
                    raise ValueError('Google Sheet 交易紀錄格式無效')
                self.merge_synced_rows(remote.get('model_trades', []))
                self.month_versions[month] = remote.get('model_version') or versions.get(month)
            with closing(self.db()) as db, db:
                # Rotate across dirty rows so frequently changing positions cannot starve later trades.
                rows = db.execute('SELECT rowid,body FROM trades WHERE month=? AND dirty=1 ORDER BY (rowid<=?),rowid LIMIT 20',
                                  (month, self.sync_positions.get(month, 0))).fetchall()
            batch, size = [], 0
            for rowid, body in rows:
                record = json.loads(body)
                if record.get('識別版本') == 2 and catalog.get('model_reentry') != 1:
                    blocked = True
                    continue
                length = len(body.encode('utf-8'))
                if batch and size + length > 256_000:
                    break
                batch.append((rowid, record))
                size += length
            if not batch:
                continue
            sent = {row['交易ID']: row for _, row in batch}
            actual = self.remote_save(scope, {'model_schema': 1, 'model_trades': list(sent.values())}, delta=True)
            saved = actual.get('model_trades', [])
            if not isinstance(saved, list):
                raise ValueError('Google Sheet 交易紀錄格式無效；保留待同步')
            saved = {r['交易ID']: r for r in merge_trades(saved, saved)}
            if any(key not in saved or merge_trades([saved[key]], [row])[0] != saved[key] for key, row in sent.items()):
                raise ValueError('交易事件回讀不一致；保留待同步並重試')
            self.merge_synced_rows(list(saved.values()), sent)
            # A concurrent device may have written between our read and POST; fetch it next cycle.
            if actual['_previous_version'] == (self.month_versions.get(month) or ''):
                self.month_versions[month] = actual['_committed_version']
            self.sync_positions[month] = batch[-1][0]
        self.hydrated = True
        self.synced_at = datetime.now(TZ).strftime('%Y/%m/%d %H:%M:%S')
        with closing(self.db()) as db, db:
            remaining = db.execute('SELECT COUNT(*) FROM trades WHERE dirty=1').fetchone()[0]
        self.status, self.error = (f'已回讀驗證；尚有 {remaining} 筆待同步' if remaining else '已回讀驗證'), ('Apps Script 尚未支援 V1.2 執行鏈；新紀錄保留本機待同步' if blocked else '')

    def run(self):
        pending = None
        busy_message = '本機資料庫忙碌；保留原批次並自動重試'
        while not self.closed:
            try:
                if pending is None:
                    pending = self.queue.get(timeout=.25)
                self.process(*pending)
                pending = None
                if self.error == busy_message or (self.error == '記錄佇列已滿；保留資料缺口，不阻塞行情' and self.queue.empty()):
                    self.error = ''
            except queue.Empty:
                pass
            except Exception as exc:
                if database_busy(exc):
                    self.error = busy_message
                    self.stop.wait(.25)
                    continue
                pending = None
                self.error = '本機記錄失敗：' + type(exc).__name__
            if time.monotonic() - self.last_gap_check >= 5:
                self.last_gap_check = time.monotonic()
                try:
                    self.mark_gaps()
                except Exception as exc:
                    self.error = '資料缺口記錄待重試：' + type(exc).__name__

    def run_sync(self):
        # Slow HTTP/retries must never stop the ordered local evidence writer.
        while not self.stop.is_set():
            if time.monotonic() - self.last_sync >= self.sync_seconds:
                self.last_sync = time.monotonic()
                try:
                    with closing(self.db()) as db, db:
                        pending = db.execute('SELECT 1 FROM trades WHERE dirty=1 LIMIT 1').fetchone()
                    if pending or time.monotonic() - self.last_activity < 120:
                        self.sync()
                    if self.archive_requests:
                        self.load_event_archive(sorted(self.archive_requests)[0])
                except Exception as exc:
                    self.error = '本機資料庫忙碌；保留待同步並自動重試' if database_busy(exc) else str(exc)
                    self.status = '待同步／本機保留'
                finally:
                    self.last_sync = time.monotonic()
            self.stop.wait(.25)

    def close(self):
        self.closed = True
        self.stop.set()
        for thread in (self.thread, self.sync_thread):
            if thread:
                thread.join(timeout=1)
