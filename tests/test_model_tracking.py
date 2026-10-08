"""Behavior checks for verified fills, cross-device merging, restart, export and UI."""
import ast
import io
import json
import shutil
import subprocess
import threading
from pathlib import Path
from unittest.mock import patch
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook
from model_tracking import DEFAULT_COSTS, ModelTracker, advance, new_trade, merge_trades, performance, excel_report, normalize_costs

ROOT = Path(__file__).parents[1]


def signal(**changes):
    return {'交易日':'2026-10-08', '市場':'股票', '商品鍵':'2330', '代碼':'2330', '名稱':'台積電',
            '策略':'當沖', '策略版本':'V1.0', '方向':'多頭', '進場價':100., '停損價':95., '目標價':110.,
            '來源時間':'2026-10-08T09:30:00+08:00', '訊號價':100., '跳動點':.5, '乘數':1,
            '成本設定':dict(DEFAULT_COSTS), '指標快照':{'VWAP':99.,'均線':98.}, '市場環境':'偏多', **changes}


def quote(price, second=0, **kw):
    return {'價格':price, '來源時間':f'2026-10-08T09:30:{second:02d}+08:00', '交易日':'2026-10-08',
            '有效':True, '更新秒數':5, **kw}


def test_verified_fills_costs_loss_first_and_ambiguity():
    waiting = new_trade(signal())
    holding = advance(waiting, quote(100))
    assert holding['狀態'] == '模擬持倉' and holding['模擬進場價'] == 100.5
    won = advance(holding, quote(111, 1))
    assert won['狀態'] == '已平倉' and won['模擬出場價'] == 110.5
    assert won['交易成本'] > 0 and 0 < won['模擬損益'] < won['毛損益']
    stopped = advance(holding, quote(94, 1))
    assert stopped['模擬損益'] < 0
    assert advance(stopped, quote(111, 2)) == stopped  # Later target cannot turn a stop into a win.
    ambiguous = advance(holding, quote(111, 1, 最高=111, 最低=94))
    assert ambiguous['資料缺口'] and ambiguous['狀態'] == '監控中斷'
    assert performance([ambiguous])[0]['有效交易數'] == 0
    expired = advance(waiting, quote(94, 1))
    assert expired['狀態'] == '訊號失效'
    not_entered = advance(waiting, quote(99, 1))
    assert not_entered['狀態'] == '等待進場'


def test_short_futures_gaps_invalid_data_and_same_timestamp_ticks():
    base = signal(市場='期貨', 商品鍵='TX:202610', 方向='偏空', 進場價=100, 停損價=105, 目標價=90,
                  跳動點=1, 乘數=200, 期交稅率=.00002)
    holding = advance(new_trade(base), quote(100, 序號=1))
    closed = advance(holding, quote(89, 序號=2))
    assert closed['狀態'] == '已平倉' and closed['模擬損益'] > 0
    missing = advance(holding, quote(99, 1, 有效=False))
    assert missing['資料缺口'] and missing['狀態'] == '資料不足'
    resumed = advance(missing, quote(89, 2))
    assert resumed['狀態'] == '已平倉' and performance([resumed])[0]['有效交易數'] == 0
    gap = advance(holding, quote(89, 40))
    assert gap['資料缺口'] and performance([gap])[0]['有效交易數'] == 0
    overflow = advance(holding, quote(89, 1, 序號=4))
    assert overflow['資料缺口']


def test_idempotent_devices_restart_and_original_plan(tmp_path):
    path = tmp_path / 'model.sqlite'
    tracker = ModelTracker(path, start=False)
    s = signal()
    tracker.process([s, signal(進場價=100)], {'股票|2330':quote(100)})
    assert len(tracker.records()) == 1
    tracker.process([signal(進場價=101, 目標價=112)], {'股票|2330':quote(101, 1)})
    assert len(tracker.records()) == 1
    restart = ModelTracker(path, start=False)
    restart.process([s], {'股票|2330':quote(111, 2)})
    records = restart.records()
    assert len(records) == 1 and records[0]['狀態'] == '已平倉'
    assert records[0]['進場價'] == 100 and records[0]['成本設定'] == DEFAULT_COSTS
    assert records[0]['資料缺口'] and performance(records)[0]['有效交易數']==0
    restart.process([s], {'股票|2330':quote(111, 3)})
    assert len(restart.records()) == 1


def test_merge_keeps_distinct_events_terminal_and_gap():
    initial = new_trade(signal())
    holding = advance(initial, quote(100))
    closed = advance(holding, quote(111, 1))
    other = new_trade(signal(商品鍵='2408'))
    merged = merge_trades([closed, other], [holding, initial])
    assert len(merged) == 2 and merged[0]['狀態'] == '已平倉'
    assert len(merged[0]['事件']) == 3
    invalid = {**holding, '資料缺口':True}
    assert not merge_trades([closed], [invalid])[0]['資料缺口']  # The connected device proves the whole completed path.
    conflict = {**closed, '模擬進場價':101.5}
    assert merge_trades([closed], [conflict])[0]['資料缺口']


def test_failed_sync_retry_remote_device_restores_and_no_overwrite(tmp_path):
    remote = {'strategy_signals':{'model_schema':1,'model_months':[], 'model_versions':{},
                                  'strategy_signal_log':[{'dedupe_key':'legacy'}]}}
    lock = threading.RLock()
    fail = [False]
    def read(scope):
        with lock:
            return json.loads(json.dumps(remote.get(scope, {})))
    def write(scope, data):
        with lock:
            if fail[0]:
                raise ConnectionError('offline')
            prior = remote.get(scope, {})
            remote[scope] = {**data, 'model_trades':merge_trades(prior.get('model_trades', []), data.get('model_trades', []))}
            month = scope.split(':')[1]
            remote['strategy_signals']['model_months'] = sorted(set(remote['strategy_signals']['model_months'] + [month]))
            remote['strategy_signals']['model_versions'][month] = str(len(remote[scope]['model_trades']))
            return read(scope)
    desktop = ModelTracker(tmp_path/'desktop.sqlite', 'fake', start=False)
    phone = ModelTracker(tmp_path/'phone.sqlite', 'fake', start=False)
    for t in (desktop, phone):
        t.remote_get, t.remote_save = read, write
    desktop.process([signal()], {'股票|2330':quote(100)})
    phone.process([signal(),signal(商品鍵='2408')], {'股票|2330':quote(100)})
    fail[0] = True
    with pytest.raises(ConnectionError):
        desktop.sync()
    with desktop.db() as db:
        assert db.execute('SELECT dirty FROM trades').fetchone()[0] == 1
    fail[0] = False
    # Stale device writes interleave; server merge must retain both instruments.
    desktop.sync(); phone.sync(); desktop.sync()
    assert len(desktop.records()) == len(phone.records()) == 2
    assert remote['strategy_signals']['strategy_signal_log'] == [{'dedupe_key':'legacy'}]
    later = ModelTracker(tmp_path/'later.sqlite', 'fake', start=False)
    later.remote_get, later.remote_save = read, write
    later.sync()
    assert len(later.records()) == 2 and later.synced_at


def test_excel_real_numeric_values_filters_and_performance():
    a = advance(advance(new_trade(signal()), quote(100)), quote(111, 1))
    b = advance(advance(new_trade(signal(商品鍵='2408')), quote(100)), quote(94, 2))
    invalid = {**a, '交易ID':'invalid', '資料缺口':True}
    metrics, groups, _ = performance([a,b,invalid])
    assert metrics['有效交易數'] == 2 and metrics['勝率(%)'] == 50
    assert metrics['最大回撤'] == -b['模擬損益']
    assert metrics['累積模擬損益'] == a['模擬損益'] + b['模擬損益']
    assert not groups.empty
    a['名稱'] = '=1+1'
    workbook = load_workbook(io.BytesIO(excel_report([a,b])))
    assert workbook.sheetnames == ['總覽','策略與指標','歷史明細','事件']
    sheet = workbook['歷史明細']
    headers = {c.value:c.column for c in sheet[1]}
    assert sheet.cell(2, headers['模擬損益']).data_type == 'n'
    assert sheet.cell(2, headers['名稱']).data_type == 's'
    assert sheet.freeze_panes == 'A2'


def test_apps_script_merge_runs_under_real_javascript():
    if not shutil.which('node'):
        pytest.skip('Node unavailable')
    program = r'''
const fs = require('fs'), vm = require('vm'), assert = require('assert');
const source = fs.readFileSync('google_apps_script.gs','utf8');
vm.runInThisContext(source);
const a={'交易ID':'A','建立時間':'09:00','來源時間':'09:00','狀態':'等待進場','成本設定':{fee:50,slip:1},'事件':[{'事件ID':'1','時間':'09:00'}]};
const b={'交易ID':'B','建立時間':'09:00','來源時間':'09:00','狀態':'等待進場'};
const closed={...a,'來源時間':'10:00','狀態':'已平倉','事件':[{'事件ID':'2','時間':'10:00'}]};
let data=mergeSignalPayload_({strategy_signal_log:[{dedupe_key:'manual','最後更新':'10:00'}],model_trades:[a]}, {model_trades:[b,closed]});
data=mergeSignalPayload_(data, {model_trades:[a]});
assert.equal(data.model_trades.length,2);
assert.equal(data.model_trades[0]['狀態'],'已平倉');
assert.equal(data.model_trades[0]['事件'].length,2);
assert(!mergeModelTrades_([closed],[{...closed,'成本設定':{slip:1,fee:50}}])[0]['資料缺口']);
assert.equal(data.strategy_signal_log[0].dedupe_key,'manual');
assert.equal(normalizeScope_('strategy_signals:202610'),'strategy_signals:202610');
assert.equal(normalizeScope_('strategy_signals:anything'),'');
assert(source.includes('lock.waitLock(30000)'));
console.log('Apps Script merge verified');
'''
    result = subprocess.run(['node','-e',program], cwd=ROOT, capture_output=True, text=True)
    assert result.returncode == 0, result.stderr


def test_new_home_tab_lazy_ui_and_excel(tmp_path):
    from streamlit.testing.v1 import AppTest
    source = (ROOT/'app.py').read_text(encoding='utf-8')
    functions = {'render_model_tracking_room','cached_model_performance','model_performance_state'}
    defs = '\n'.join(ast.unparse(n) for n in ast.parse(source).body if isinstance(n,ast.FunctionDef) and n.name in functions)
    tracker = ModelTracker(tmp_path/'ui.sqlite',start=False)
    script = '''
import streamlit as st
import pandas as pd, json, pytz, threading, time
from datetime import date,datetime,timedelta
from model_tracking import ModelTracker,DEFAULT_COSTS,performance,excel_report,normalize_costs
'''+f"tracker=ModelTracker({str(tmp_path/'ui.sqlite')!r},start=False)\n"+'''
get_model_tracker=lambda *a: tracker
get_app_secret=lambda *a: None
load_config=lambda: {}
render_strategy_validation_room=lambda: st.info('legacy')
'''+defs+'''
outer,other=st.tabs(['交易損益室','其他'],key='main',on_change='rerun')
with outer:
    first,second=st.tabs(['模型勝率追蹤','當沖損益室'],default='模型勝率追蹤',key='profit',on_change='rerun')
    with first:
        if outer.open and first.open:
            render_model_tracking_room()
'''
    app = AppTest.from_string(script).run()
    assert not app.exception and app.session_state['profit'] == '模型勝率追蹤'
    assert len(app.get('download_button')) == 1
    app.text_input(key='model_code').set_value('2330').run()
    assert not app.exception
    tracker.close()


def test_invalid_settings_and_plans_fail_safely():
    assert normalize_costs({'stock_shares':-2,'stock_discount':float('nan')}) == DEFAULT_COSTS
    with pytest.raises(ValueError):
        new_trade(signal(停損價=110))


def test_shared_writer_two_live_streams_and_failover(tmp_path):
    tracker=ModelTracker(tmp_path/'shared.sqlite',start=False)
    tick1=quote(100,0,串流識別='desktop:0',序號=1)
    tracker.process([signal()], {'股票|2330':{**tick1,'成交證據':[tick1]}})
    phone1=quote(100,0,串流識別='phone:0',序號=1)
    phone2=quote(101,1,串流識別='phone:0',序號=2)
    tracker.process([signal()],{'股票|2330':{**phone2,'成交證據':[phone1,phone2]}})
    assert len(tracker.records()) == 1 and not tracker.records()[0]['資料缺口']
    down=quote(100,2,有效=False,串流識別='desktop:0')
    tracker.process([],{'股票|2330':down})
    phone3=quote(111,3,串流識別='phone:0',序號=3)
    tracker.process([],{'股票|2330':{**phone3,'成交證據':[phone1,phone2,phone3]}})
    row=tracker.records()[0]
    assert row['狀態']=='已平倉' and not row['資料缺口']
    assert performance([row])[0]['有效交易數']==1


def test_full_app_profit_home_is_lazy_and_preserves_other_tabs(tmp_path):
    from streamlit.testing.v1 import AppTest
    import requests
    source=(ROOT/'app.py').read_text(encoding='utf-8')
    anchor='tab1, tab_fibo, tab2, tab_db, tab_company, tab3 = st.tabs(['
    setup = f"CONFIG_FILE={str(tmp_path/'config.json')!r}\n_ui_tracker=ModelTracker({str(tmp_path/'model.sqlite')!r},start=False)\nget_model_tracker=lambda *a: _ui_tracker\n" + """
get_app_secret=lambda key,default=None: default
render_postclose_maintenance=lambda: None
st.session_state.setdefault('main_workspace_active_tab','💰 交易損益室 💰')
"""
    source=source.replace(anchor,setup+anchor)
    with patch('requests.get',side_effect=requests.ConnectionError('offline check')), \
         patch('requests.post',side_effect=requests.ConnectionError('offline check')), \
         patch('yfinance.download',return_value=pd.DataFrame()), \
         patch('yfinance.Ticker',return_value=SimpleNamespace(history=lambda *a,**kw:pd.DataFrame(),fast_info={},info={})):
        app=AppTest.from_string(source,default_timeout=120).run()
        assert not app.exception
        assert app.session_state['profit_room_active_tab']=='模型勝率追蹤'
        assert len(app.get('download_button')) >= 1
        app.session_state['main_workspace_active_tab']='⚡ 股期戰略室 ⚡'
        app.run()
        assert not app.exception
        assert not any('資料狀態：' in c.value and '同步時間：' in c.value for c in app.caption)


def test_apps_script_full_read_write_lock_and_partial_commit_recovery():
    if not shutil.which('node'):
        pytest.skip('Node unavailable')
    program=r"""
const fs=require('fs'),vm=require('vm'),assert=require('assert');
const rows=[['scope','part','data','updated_at']];
let held=false,fail=false;
const sheet={getLastRow:()=>rows.length,
  getRange:(start,col,count,width)=>({
    getDisplayValues:()=>Array.from({length:count},(_,i)=>Array.from({length:width},(_,j)=>String((rows[start+i-1]||[])[col+j-1]||''))),
    setNumberFormat(){return this;},
    setValues(values){assert(held); for(let i=0;i<values.length;i++) {rows[start+i-1]=values[i]; if(fail && values[i][0].includes(':')) {fail=false; throw Error('partial write');}} return this;}
  }),deleteRow:(r)=>rows.splice(r-1,1),clearContents:()=>rows.splice(0)};
global.SpreadsheetApp={getActiveSpreadsheet:()=>({getSheetByName:()=>sheet}),flush:()=>{}};
global.LockService={getScriptLock:()=>({waitLock:()=>{assert(!held);held=true;},releaseLock:()=>held=false})};
global.ContentService={MimeType:{JSON:'json'},createTextOutput:(body)=>({setMimeType:()=>JSON.parse(body)})};
vm.runInThisContext(fs.readFileSync('google_apps_script.gs','utf8'));
function post(scope,data){return doPost({postData:{contents:JSON.stringify({scope,data,updated_at:new Date().toISOString()})}});}
const a={'交易ID':'a'.repeat(64),'交易日':'2026-10-08','市場':'股票','商品鍵':'2330','策略':'當沖','策略版本':'V1.0','方向':'多頭','進場價':100,'停損價':95,'目標價':110,'建立時間':'2026-10-08T09:00:00+08:00','來源時間':'2026-10-08T09:00:00+08:00','狀態':'等待進場','事件':[]};
assert(post('strategy_signals',{strategy_signal_log:[{dedupe_key:'manual'}]}).success);
assert(post('strategy_signals:202610',{model_trades:[a]}).success);
const stale={model_trades:[{...a,'交易ID':'b'.repeat(64)}]};
assert(post('strategy_signals:202610',stale).success);
let read=doGet({parameter:{scope:'strategy_signals:202610'}});
assert.equal(read.data.model_trades.length,2);
let catalog=doGet({parameter:{scope:'strategy_signals'}}).data;
assert(catalog.model_months.includes('202610') && catalog.strategy_signal_log[0].dedupe_key==='manual');
fail=true;
assert(!post('strategy_signals:202610',{model_trades:[{...a,'交易ID':'c'.repeat(64),'指標快照':{text:'x'.repeat(60000)}}]}).success);
read=doGet({parameter:{scope:'strategy_signals:202610'}});
assert.equal(read.data.model_trades.length,2); // Last complete generation survives partial write.
assert(post('strategy_signals:202610',{model_trades:[{...a,'交易ID':'c'.repeat(64),'指標快照':{text:'x'.repeat(60000)}}]}).success);
assert.equal(doGet({parameter:{scope:'strategy_signals:202610'}}).data.model_trades.length,3);
assert(!post('strategy_signals:202609',{model_trades:[a]}).success);
assert(!post('strategy_signals:202610',{model_trades:[{...a,'交易ID':'__proto__'}]}).success);
assert(!held);
"""
    result=subprocess.run(['node','-e',program],cwd=ROOT,capture_output=True,text=True)
    assert result.returncode==0,result.stderr


def real_stock_adapter():
    import importlib.util
    from collections import deque
    spec=importlib.util.spec_from_file_location('core_checks',ROOT/'tests/test_core_calculations.py')
    core=importlib.util.module_from_spec(spec);spec.loader.exec_module(core)
    names=('model_tracking_samples','_safe_number','_as_float','get_taiwan_tick_size','round_to_tick',
           'get_tick_size','move_tick','fmt_price','_format_compact_number','determine_stock_direction',
           'calculate_risk_filter_result','calculate_daytrade_filter_result','build_trade_plan',
           'classify_signal_state','parse_trade_plan_numbers','parse_strategy_data_time','build_data_health',
           'calculate_entry_confidence','calculate_market_alignment','market_risk_checked_for_row')
    ns=core.load_app_symbols(*names)
    ns.update(deque=deque,DEFAULT_COSTS=DEFAULT_COSTS,MODEL_VERSION='V1.0',model_market_open=lambda r:True)
    return ns


def stock_rows_with_evidence(ns, count=6):
    from collections import deque
    now=ns['datetime'].now(ns['pytz'].timezone('Asia/Taipei'))
    stamp=now.isoformat()
    rows=[]
    buffers={}
    for i in range(count):
        code=str(2330+i)
        rows.append({'代號':code,'名稱':'驗證股票','收盤價':100.5,'_quote_time':stamp,'_ma5':98.,
                     '_risk_ma20':98.,'_risk_ma20_slope':.5,'_risk_atr14':4.,'_risk_close_position':80.,
                     '_risk_prev_high':99.,'_risk_prev_low':96.,'_daytrade_vwap':99.,
                     '_daytrade_or_high':100.,'_daytrade_or_low':98.,'_daytrade_close':100.5,
                     '_daytrade_volume_ratio':2.,'_daytrade_data_time':stamp,'_daytrade_phase':'開盤區間完成',
                     '當日漲停價':110.,'當日跌停價':90.})
        buffers[code]=deque([{'來源時間':stamp,'價格':100.5,'序號':1,'串流識別':'desktop:0','有效':True}],maxlen=4096)
    state={'lock':threading.RLock(),'model_tick_buffers':buffers}
    ns['_stream_state']=lambda api:state
    config={'risk':{'updated':stamp},'costs':dict(DEFAULT_COSTS),'market_bias':'偏多','min_score':75,
            'extension':2.,'block_attention':True,'rules_version':'V1.0-test'}
    return pd.DataFrame(rows), config


def test_actual_stock_strategy_adapter_records_without_quote_queries(tmp_path):
    ns=real_stock_adapter()
    rows,config=stock_rows_with_evidence(ns)
    original=rows.copy(deep=True)
    records,quotes=ns['model_tracking_samples'](rows,True,'當沖','多',config,5,object())
    assert len(records)==len(quotes)==6
    assert all(r['策略']=='當沖' and r['進場價']==100.5 and r['指標快照']['ATR']==4 for r in records)
    pd.testing.assert_frame_equal(rows,original)
    t=ModelTracker(tmp_path/'integration.sqlite',start=False)
    t.process(records,quotes)
    assert len(t.records())==6


def benchmark_tracking(path):
    import time, statistics, tracemalloc
    ns=real_stock_adapter();rows,config=stock_rows_with_evidence(ns)
    t=ModelTracker(path,start=False)
    timings={'before':[],'after':[]}
    tracemalloc.start()
    for key in ('before','after'):
        cpu=time.process_time();started=time.perf_counter()
        for _ in range(200):
            begin=time.perf_counter()
            copied=rows.copy(deep=True)
            if key=='after':
                signals,quotes=ns['model_tracking_samples'](copied,True,'當沖','多',config,5,object())
                t.submit(signals,quotes)
                t.queue.get_nowait()  # Isolate nonblocking table-side overhead; writer measured separately.
            timings[key].append(time.perf_counter()-begin)
        timings[key+'_cpu']=time.process_time()-cpu
        timings[key+'_wall']=time.perf_counter()-started
    signals,quotes=ns['model_tracking_samples'](rows,True,'當沖','多',config,5,object())
    writer_start=time.perf_counter();t.process(signals,quotes);writer=time.perf_counter()-writer_start
    _,peak=tracemalloc.get_traced_memory();tracemalloc.stop()
    base=advance(advance(new_trade(signal()),quote(100)),quote(111,1))
    with t.db() as db:
        for i in range(10000):
            r={**base,'交易ID':f'history-{i}'}
            t.store(db,r,0)
    start=time.perf_counter();t.process(signals,quotes);writer_large=time.perf_counter()-start
    start=time.perf_counter();history=t.records('202610','202610');load=time.perf_counter()-start
    start=time.perf_counter();performance(history);stats=time.perf_counter()-start
    tracemalloc.start()
    from collections import deque
    buffers=[deque(({'價格':100+i/100,'來源時間':f'2026-10-08T09:30:00.{i:06d}+08:00','序號':i+1,'串流識別':'api:0:1','有效':True} for i in range(4096)),maxlen=4096) for _ in range(11)]
    buffer_peak=tracemalloc.get_traced_memory()[1];tracemalloc.stop()
    assert len(buffers)==11
    return {'targets':6,'iterations':200,'before_table_ms':statistics.median(timings['before'])*1000,
            'after_table_ms':statistics.median(timings['after'])*1000,'writer_ms':writer*1000,
            'writer_10000_history_ms':writer_large*1000,'history_10000_load_ms':load*1000,
            'history_10000_stats_ms':stats*1000,'table_before_cpu_s':timings['before_cpu'],
            'table_after_cpu_s':timings['after_cpu'],'peak_allocated_bytes':peak,'eleven_full_tick_buffers_bytes':buffer_peak,'extra_broker_queries':0}


def test_actual_futures_adapter_uses_index_point_value_and_existing_thresholds():
    from datetime import timedelta
    from collections import deque
    import importlib.util
    spec=importlib.util.spec_from_file_location('core_checks',ROOT/'tests/test_core_calculations.py')
    core=importlib.util.module_from_spec(spec);spec.loader.exec_module(core)
    ns=core.load_app_symbols('model_tracking_samples','enrich_futures_strategy_rows','_safe_number',
                            'get_futures_tick_size','FUTURES_FIXED_TICK_SIZES','parse_strategy_data_time',
                            'build_data_health','calculate_market_alignment','parse_trade_plan_numbers','calculate_entry_confidence')
    ns.update(deque=deque,DEFAULT_COSTS=DEFAULT_COSTS,MODEL_VERSION='V1.0',model_market_open=lambda r:True)
    now=ns['datetime'].now(ns['pytz'].timezone('Asia/Taipei'));ts=now.isoformat()
    ns['get_futures_trading_date']=lambda dt:dt
    ns['futures_expiry_date']=lambda month:now.date()+timedelta(days=30)
    ns['resolve_shioaji_futures_contract']=lambda *a:SimpleNamespace(code='TX202610')
    state={'lock':threading.RLock(),'model_tick_buffers':{'TX202610':deque([{'來源時間':ts,'價格':100.,'序號':1,'有效':True}])}}
    ns['_stream_state']=lambda api:state
    rows=pd.DataFrame([{'契約鍵':'TX:202610','期貨代碼':'TX','契約月份':'202610','商品類型':'指數','乘數':1,
                        '名稱':'臺指期貨','收盤價':100.,'當日成交口數':2000,'未平倉量':500,'買價':99.,'賣價':100.,
                        '報價時間':ts,'進出場點位':'進 100｜停 95｜目 110','方向':'偏多','VWAP':99.,'ATR':4.}])
    signals,quotes=ns['model_tracking_samples'](rows,False,'當沖','自動',{'market_bias':'偏多','costs':DEFAULT_COSTS},5,object())
    assert len(signals)==1 and signals[0]['乘數']==200 and signals[0]['期交稅率']==.00002
    assert len(quotes)==1


def test_new_day_daytrade_does_not_hold_overnight_or_block_next_signal(tmp_path):
    t=ModelTracker(tmp_path/'newday.sqlite',start=False)
    t.process([signal()],{'股票|2330':quote(100)})
    nextday=signal(交易日='2026-10-09',來源時間='2026-10-09T09:30:00+08:00')
    q={**quote(100),'交易日':'2026-10-09','來源時間':nextday['來源時間']}
    t.process([nextday],{'股票|2330':q})
    records=t.records()
    assert len(records)==2
    assert next(r for r in records if r['交易日']=='2026-10-08')['狀態']=='訊號失效'
    assert performance(records)[0]['有效交易數']==0


def test_wave_continuous_stream_survives_market_rest():
    first=signal(策略='波段',來源時間='2026-10-08T13:30:00+08:00')
    r=advance(new_trade(first),{'價格':100,'來源時間':first['來源時間'],'交易日':'2026-10-08','有效':True,'序號':1,'串流識別':'api:0'})
    nextquote={'價格':111,'來源時間':'2026-10-09T09:00:00+08:00','交易日':'2026-10-09','有效':True,
               '序號':2,'串流識別':'api:0','休市銜接':True}
    closed=advance(r,nextquote)
    assert closed['狀態']=='已平倉' and not closed['資料缺口']


def test_incremental_statistics_and_remote_correction():
    a=advance(advance(new_trade(signal()),quote(100)),quote(111,1))
    b=advance(advance(new_trade(signal(商品鍵='2408')),quote(100)),quote(94,2))
    state={}
    assert performance([a],state)[0]['有效交易數']==1 and state['added']==1
    assert performance([a],state)[0]['有效交易數']==1 and state['added']==0
    assert performance([a,b],state)[0]['有效交易數']==2 and state['added']==1
    assert performance([{**a,'資料缺口':True},b],state)[0]['有效交易數']==1
    assert state['added']==1 and performance([b])[0]==performance([b],state)[0]


def test_verified_close_after_session_end_and_cost_conflict(tmp_path):
    t=ModelTracker(tmp_path/'close.sqlite',start=False)
    s=signal(來源時間='2026-10-08T13:29:59+08:00')
    q={**quote(100),'來源時間':s['來源時間']}
    t.process([s],{'股票|2330':q})
    q={**quote(101),'來源時間':'2026-10-08T13:30:00+08:00','休市':True,'收盤確認':True}
    t.process([],{'股票|2330':q})
    row=t.records()[0]
    assert row['狀態']=='已平倉' and row['出場原因']=='收盤' and not row['資料缺口']
    other={**row,'成本設定':{**DEFAULT_COSTS,'slippage_ticks':0}}
    assert merge_trades([row],[other])[0]['資料缺口']
    invalid={**row,'資料缺口':True,'結案時間':'2026-10-08T13:29:59+08:00'}
    assert not merge_trades([invalid],[row])[0]['資料缺口']


def test_new_connection_requires_overlapping_original_evidence():
    first=advance(new_trade(signal()),quote(100,串流識別='before:0',序號=1))
    resumed=advance(first,quote(111,1,串流識別='after:0',序號=1))
    assert resumed['資料缺口'] and performance([resumed])[0]['有效交易數']==0


def test_concurrent_devices_same_signal_and_different_symbols(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    remote={'strategy_signals':{'model_schema':1,'model_months':[],'model_versions':{}}}
    lock=threading.RLock(); barrier=threading.Barrier(2); seen=set()
    def read(scope):
        identity=threading.get_ident()
        with lock:
            result=json.loads(json.dumps(remote.get(scope,{})))
            first=':' in scope and identity not in seen
            if first: seen.add(identity)
        if first: barrier.wait(timeout=5)
        return result
    def write(scope,data):
        with lock:
            remote[scope]={'model_schema':1,'model_trades':merge_trades(remote.get(scope,{}).get('model_trades',[]),data['model_trades'])}
            remote['strategy_signals']['model_months']=['202610']
            remote['strategy_signals']['model_versions']['202610']=str(len(remote[scope]['model_trades']))
        return read(scope)
    trackers=[ModelTracker(tmp_path/f'device-{i}.sqlite','fake',start=False) for i in range(2)]
    for i,t in enumerate(trackers):
        t.remote_get,t.remote_save=read,write
        t.process([signal(),signal(商品鍵=str(2408+i))],{'股票|2330':quote(100)})
    with ThreadPoolExecutor(max_workers=2) as executor:
        list(executor.map(lambda t:t.sync(),trackers))
    for t in trackers:
        t.remote_get=lambda scope:json.loads(json.dumps(remote.get(scope,{})))
        t.sync()
        assert len(t.records())==3
    assert len({r['交易ID'] for r in remote['strategy_signals:202610']['model_trades']})==3


def test_old_backend_preserves_pending_without_writing(tmp_path):
    tracker=ModelTracker(tmp_path/'old-backend.sqlite','fake',start=False)
    tracker.process([signal()],{'股票|2330':quote(100)})
    tracker.remote_get=lambda scope: {'strategy_signal_log':[]}
    tracker.remote_save=lambda *a:pytest.fail('Old backend must not receive model writes')
    with pytest.raises(ValueError,match='尚未部署'):
        tracker.sync()
    with tracker.db() as db:
        assert db.execute('SELECT dirty FROM trades').fetchone()[0]==1


def test_stale_tick_is_not_a_fresh_snapshot_signal():
    from datetime import timedelta
    ns=real_stock_adapter();rows,config=stock_rows_with_evidence(ns,1)
    state=ns['_stream_state'](None)
    state['model_tick_buffers']['2330'][0]['來源時間']=(ns['datetime'].now(ns['pytz'].timezone('Asia/Taipei'))-timedelta(minutes=2)).isoformat()
    signals,quotes=ns['model_tracking_samples'](rows,True,'當沖','多',config,5,object())
    assert not signals and not quotes['股票|2330']['有效']
