"""V1.2 recording changes preserve execution evidence and legacy results."""
import io
import json
import shutil
import subprocess
from pathlib import Path
import pytest
from openpyxl import load_workbook
from model_tracking import ModelTracker, advance, new_trade, trade_id, merge_trades, performance, plan_performance, excel_report, event
from test_model_tracking import signal, quote


def test_same_timestamp_replay_cannot_rewind_sequence_or_create_false_gap():
    row = advance(advance(new_trade(signal()), quote(99.5, 序號=1, 串流識別='api:0')),
                  quote(100, 序號=2, 串流識別='api:0'))
    original = dict(row)
    row = advance(row, quote(99.5, 序號=1, 串流識別='api:0'))
    assert row == original
    row = advance(row, quote(101, 1, 序號=3, 串流識別='api:0'))
    assert not row['資料缺口'] and row['最後序號'] == 3
    assert advance(row, quote(102, 2, 序號=5, 串流識別='api:0'))['資料缺口']


def test_invalid_replayed_tick_keeps_real_gap_but_cannot_rewind_sequence():
    row = advance(new_trade(signal()), quote(100, 序號=2, 串流識別='api:0'))
    invalid = advance(row, quote(99, 1, 序號=1, 串流識別='api:0', 有效=False))
    assert invalid['資料缺口'] and invalid['最後序號'] == 2
    recovered = advance(invalid, quote(101, 2, 序號=3, 串流識別='api:0'))
    assert recovered['資料缺口'] and not any(e.get('缺漏證據') for e in recovered['事件'])


def test_source_metadata_does_not_change_legacy_identity_or_results():
    stock = signal()
    row = new_trade(stock)
    assert row['計畫來源'] == '股期戰略室' and row['計畫類型'] == '股票當沖'
    assert row['交易ID'] == trade_id(stock)
    legacy = {key: value for key, value in row.items() if key not in ('計畫來源', '計畫類型')}
    assert merge_trades([legacy], [row])[0] == row


def test_waiting_reason_survives_entry_and_invalidation_without_repeated_events():
    row = advance(new_trade(signal()), quote(99, 序號=1, 串流識別='api:0'))
    assert row['未進場分類'] == '未觸價'
    count = len(row['事件'])
    row = advance(row, quote(99, 1, 序號=2, 串流識別='api:0'))
    assert len(row['事件']) == count
    held = advance(row, quote(100, 2, 序號=3, 串流識別='api:0'))
    assert held['最後進場等待原因'] == '尚未觸及原進場價'
    invalid = advance(row, quote(94, 2, 序號=3, 串流識別='api:0'))
    assert invalid['訊號失效原因'] and invalid['最後進場等待原因']


def v12_quote(price, second, **changes):
    tick = quote(price, second, 序號=second + 1, 串流識別='api:0')
    return {**tick, '策略': '當沖', '策略版本': 'V1.2-test', '成交證據': [tick], **changes}


@pytest.mark.parametrize('first_exit', [94, 111])
def test_reentry_requires_fresh_reset_and_uses_independent_execution(tmp_path, first_exit):
    tracker = ModelTracker(tmp_path/'reentry.sqlite', start=False)
    initial = signal(策略版本='V1.2-test')
    tracker.process([initial], {'股票|2330': v12_quote(100, 0)})
    tracker.process([], {'股票|2330': v12_quote(first_exit, 1)})
    tracker.process([], {'股票|2330': v12_quote(111, 2)})
    assert len(tracker.records()) == 1
    tracker.process([], {'股票|2330': v12_quote(99, 3)})
    tracker = ModelTracker(tracker.path, start=False)
    tracker.process([], {'股票|2330': v12_quote(100, 4)})
    rows = sorted(tracker.records(), key=lambda row: row['進場次數'])
    assert len(rows) == 2 and rows[1]['狀態'] == '模擬持倉'
    assert rows[0]['原始計畫ID'] == rows[1]['原始計畫ID']
    assert rows[0]['交易ID'] != rows[1]['交易ID'] and rows[1]['前次執行ID'] == rows[0]['交易ID']
    assert rows[1]['進場時間'] > rows[0]['結案時間']
    assert not rows[1].get('結案時間') and not rows[1]['資料缺口']
    tracker.process([], {'股票|2330': v12_quote(94, 5)})
    assert len(tracker.records()) == 2 and all(row['狀態'] == '已平倉' for row in tracker.records())


def test_one_buffer_records_two_cycles_and_replay_is_idempotent(tmp_path):
    tracker = ModelTracker(tmp_path/'buffer.sqlite', start=False)
    evidence = [v12_quote(price, i)['成交證據'][0] for i, price in enumerate((100, 111, 99, 100, 94))]
    latest = v12_quote(94, 4, 成交證據=evidence)
    tracker.process([signal(策略版本='V1.2-test')], {'股票|2330': latest})
    records = tracker.records()
    assert len(records) == 2 and all(row['狀態'] == '已平倉' for row in records)
    tracker.process([signal(策略版本='V1.2-test')], {'股票|2330': latest})
    assert tracker.records() == records


def test_two_cycles_use_published_permission_instead_of_latest_wait_condition(tmp_path):
    tracker = ModelTracker(tmp_path/'permissions.sqlite', start=False)
    initial = v12_quote(99.5, 0, 進場品質許可=True, 進場品質方向='多頭', 進場品質時間=quote(0)['來源時間'])
    tracker.process([signal(策略版本='V1.2-test', 進場品質檢查=True)], {'股票|2330': initial})
    evidence = [v12_quote(price, i)['成交證據'][0] for i, price in enumerate((99.5, 100, 111, 99, 100, 94))]
    latest = v12_quote(94, 5, 成交證據=evidence, 進場品質許可=False, 進場品質方向='多頭', 進場品質時間=quote(0, 5)['來源時間'])
    tracker.process([], {'股票|2330': latest})
    rows = tracker.records()
    assert len(rows) == 2 and all(row['狀態'] == '已平倉' and not row['資料缺口'] for row in rows)


@pytest.mark.parametrize('kind', ['bc', 'bp', 'spread', 'index'])
def test_operation_reentry_uses_new_contract_evidence_and_fresh_costs(tmp_path, kind):
    from test_operation_model_tracking import adapter, spec, option, tick, sample
    ns, state = adapter()
    put, spread = kind == 'bp', kind == 'spread'
    plan = spec(ns, None if kind == 'index' else option(spread=spread, put=put), spread=spread,
                direction='偏空' if put else '偏多')
    plan['策略版本'] = 'V1.2-operation-test'
    tracker = ModelTracker(tmp_path/(kind+'.sqlite'), start=False)
    for second, underlying in enumerate((19999, 20000, 20200, 19999, 20000, 20200)):
        if put:
            underlying = 40000 - underlying
        if spread:
            tick(state, 'P1', (31, 30, 10, 31, 25, 10)[second], second)
            tick(state, 'P2', (13, 12, 4, 13, 10, 4)[second], second)
        elif kind != 'index':
            tick(state, 'P1' if put else 'C1', (30, 32, 47, 31, 35, 48)[second], second)
        tick(state, 'TMFJ6', underlying, second)
        tracker.process(*sample(ns, [plan], second))
    rows = sorted(tracker.records(), key=lambda row: row['進場次數'])
    assert len(rows) == 2 and all(row['狀態'] == '已平倉' and not row['資料缺口'] for row in rows)
    assert rows[1]['交易成本'] > 0 and rows[1]['進場時間'] > rows[0]['結案時間']
    if kind != 'index':
        assert rows[1]['進場權利金'] != rows[0]['進場權利金']
        assert set(rows[1]['成交證據識別']['權利金證據']) == {leg['代碼'] for leg in rows[1]['選擇權腿']}
        if shutil.which('node'):
            script = '''const assert=require('node:assert/strict'),crypto=require('node:crypto');
global.Utilities={DigestAlgorithm:{SHA_256:'sha256'},Charset:{UTF_8:'utf8'},computeDigest:(a,t)=>Array.from(crypto.createHash(a).update(t).digest())};
''' + Path('google_apps_script.gs').read_text(encoding='utf-8') + '\nconst rows=' + json.dumps(rows, ensure_ascii=False) + ''';
validateModelExecutions_(rows,rows);assert(rows.every(r=>r['同步確認']));
const stale=JSON.parse(JSON.stringify(rows[1]));Object.values(stale['成交證據識別']['權利金證據']).forEach(p=>p['來源時間']='2000-01-01T00:00:00+08:00');
assert.throws(()=>validateModelExecutions_(rows,[stale]));
'''
            result = subprocess.run(['node'], input=script, text=True, encoding='utf-8', capture_output=True)
            assert result.returncode == 0, result.stderr


def test_invalid_quote_cancels_pending_reentry_reset(tmp_path):
    tracker = ModelTracker(tmp_path/'invalid.sqlite', start=False)
    tracker.process([signal(策略版本='V1.2-test')], {'股票|2330': v12_quote(100, 0)})
    for second, price in enumerate((111, 99), 1):
        tracker.process([], {'股票|2330': v12_quote(price, second)})
    tracker.process([], {'股票|2330': v12_quote(99, 3, 有效=False)})
    tracker.process([], {'股票|2330': v12_quote(100, 4)})
    assert len(tracker.records()) == 1


def test_source_and_contract_legs_participate_in_v12_plan_identity():
    base = signal(策略版本='V1.2-test', 市場='期貨')
    source = {**base, '計畫來源': '指數操作計畫', '計畫類型': '日線／15分K'}
    assert trade_id(base) != trade_id(source)
    assert trade_id({**source, '計畫類型': '短波5分K'}) != trade_id(source)


def test_same_time_new_transport_is_not_silently_treated_as_duplicate():
    row = advance(new_trade(signal()), quote(100, 序號=1, 串流識別='before:0'))
    changed = advance(row, quote(100, 序號=1, 串流識別='after:0'))
    assert changed['資料缺口'] and changed['串流識別'] == 'after:0'


def test_plan_metrics_distinguish_first_and_repeated_pending_executions(tmp_path):
    tracker = ModelTracker(tmp_path/'metrics.sqlite', start=False)
    evidence = [v12_quote(price, i)['成交證據'][0] for i, price in enumerate((100, 111, 99, 100, 94))]
    tracker.process([signal(策略版本='V1.2-test')], {'股票|2330': v12_quote(94, 4, 成交證據=evidence)})
    rows = tracker.records()
    metrics, _, _ = performance(rows)
    plans, grouped = plan_performance(rows)
    assert metrics['有效交易數'] == 2 and metrics['勝率(%)'] == 50
    assert plans['原始計畫數'] == 1 and plans['首次進場勝率(%)'] == 100 and plans['二次以上勝率(%)'] == 0
    assert grouped.iloc[0]['計畫累積損益'] == metrics['累積模擬損益']
    pending = [{**row, '同步確認': False} for row in rows]
    assert performance(pending)[0]['有效交易數'] == 0


def test_excel_has_event_rows_and_lossless_oversized_text_chunks():
    row = new_trade(signal())
    text = '=1+' + '😀完整證據' * 10000
    row['事件'].append({'事件ID':'long','時間':row['來源時間'],'狀態':'資料不足','原因':text})
    workbook = load_workbook(io.BytesIO(excel_report([row])))
    headers = [cell.value for cell in workbook['歷史明細'][1]]
    assert '事件' not in headers and '事件數' in headers
    assert workbook['事件'].max_row == 3
    chunks = list(workbook['長文字'].values)
    assert ''.join(value[-1] for value in chunks[1:]) == text
    assert all(len(value[-1].encode('utf-16-le')) <= 64000 for value in chunks[1:])
    assert all(cell.data_type != 'f' for sheet in workbook for values in sheet for cell in values)


def test_interval_diagnostics_merge_without_losing_latest_end_or_true_gap():
    row = new_trade(signal(策略版本='V1.2-test'))
    row['資料缺口'] = True
    event(row, '資料不足', quote(0, 1)['來源時間'], '連線中斷')
    earlier = json.loads(json.dumps(row))
    event(row, '資料不足', quote(0, 2)['來源時間'], '連線中斷')
    assert len(row['事件']) == 2 and row['事件'][-1]['次數'] == 2
    assert row['事件'][-1]['結束時間'] == quote(0, 2)['來源時間']
    for a, b in ((earlier, row), (row, earlier)):
        merged = merge_trades([a], [b])[0]
        assert merged['事件'][-1]['次數'] == 2 and merged['資料缺口']


def test_old_months_are_loaded_only_on_request_except_pending_or_active(tmp_path):
    tracker = ModelTracker(tmp_path/'history.sqlite', start=False)
    catalog = {'model_history':1, 'model_months':['202001','202609','202610'], 'model_active_months':{'202001':False}}
    assert '202001' not in tracker.sync_months(catalog, [])
    assert '202001' in tracker.sync_months(catalog, ['202001'])
    tracker.request_history('2020-01-01', '2020-01-31')
    assert '202001' in tracker.sync_months(catalog, [])


def test_apps_script_checks_shared_execution_chain_and_option_evidence(tmp_path):
    node = shutil.which('node')
    if not node:
        pytest.skip('Node required for real Apps Script validation')
    tracker = ModelTracker(tmp_path/'chain.sqlite', start=False)
    evidence = [v12_quote(price, i)['成交證據'][0] for i, price in enumerate((100, 111, 99, 100, 94))]
    tracker.process([signal(策略版本='V1.2-test')], {'股票|2330': v12_quote(94, 4, 成交證據=evidence)})
    rows = sorted(tracker.records(), key=lambda row: row['進場次數'])
    script = '''const assert=require('node:assert/strict'),crypto=require('node:crypto');
global.Utilities={DigestAlgorithm:{SHA_256:'sha256'},Charset:{UTF_8:'utf8'},computeDigest:(a,t)=>Array.from(crypto.createHash(a).update(t).digest())};
''' + Path('google_apps_script.gs').read_text(encoding='utf-8') + '\nconst rows=' + json.dumps(rows, ensure_ascii=False) + ''';
validateModelExecutions_(rows,rows);
assert(rows.every(r=>r['同步確認']));
const duplicated=mergeModelTrades_(rows,rows);
assert.equal(duplicated.length,2);
assert.throws(()=>validateModelExecutions_(rows,[{...rows[1],'進場次數':3}]));
assert.throws(()=>validateModelExecutions_([{...rows[0],'狀態':'模擬持倉'},rows[1]],[rows[1]]));
assert.throws(()=>validateModelExecutions_(rows,[{...rows[1],'原始計畫ID':'0'.repeat(64)}]));
'''
    result = subprocess.run([node], input=script, text=True, encoding='utf-8', capture_output=True)
    assert result.returncode == 0, result.stderr


def test_archive_roundtrip_partial_write_and_stale_writer_keep_all_evidence(tmp_path):
    node = shutil.which('node')
    if not node:
        pytest.skip('Node required for Apps Script storage checks')
    row = new_trade(signal())
    row['資料缺口'] = True
    row['事件'] += [{'事件ID':str(i), '時間':row['來源時間'], '狀態':'監控中斷', '原因':'連線中斷', '證據':'證'*200} for i in range(300)]
    row['事件'] += [{'事件ID':'recover'+str(i), '時間':row['來源時間'], '狀態':'模擬持倉', '原因':'恢復追蹤；缺口仍保留並排除有效績效'} for i in range(300)]
    row['事件'] += [{'事件ID':'entry', '時間':row['來源時間'], '狀態':'模擬持倉', '原因':'已進場'},
                    {'事件ID':'protect', '時間':row['來源時間'], '狀態':'模擬持倉', '原因':'保利更新', '保利價':101}]
    program = '''const assert=require('node:assert/strict'),crypto=require('node:crypto');
global.Utilities={DigestAlgorithm:{SHA_256:'sha256'},Charset:{UTF_8:'utf8'},computeDigest:(a,t)=>Array.from(crypto.createHash(a).update(t).digest())};
const values=[['scope','part','data','updated_at']];let held=false,fail=false;
const sheet={getLastRow:()=>values.length,getRange:(r,c,n,w)=>({
getDisplayValues:()=>Array.from({length:n},(_,i)=>Array.from({length:w},(_,j)=>String((values[r+i-1]||[])[c+j-1]||''))),
setNumberFormat(){return this;},setValues(rows){assert(held);for(let i=0;i<rows.length;i++){values[r+i-1]=rows[i];if(fail && rows[i][0].startsWith('strategy_events:')){fail=false;throw Error('partial cold write');}}return this;}
}),deleteRows:(r,n)=>{assert(held);values.splice(r-1,n);}};
global.SpreadsheetApp={getActiveSpreadsheet:()=>({getSheetByName:()=>sheet}),flush:()=>{}};
global.LockService={getScriptLock:()=>({waitLock:()=>{assert(!held);held=true;},releaseLock:()=>held=false})};
global.ContentService={MimeType:{JSON:'json'},createTextOutput:t=>({setMimeType:()=>JSON.parse(t)})};
''' + Path('google_apps_script.gs').read_text(encoding='utf-8') + '\nconst original=' + json.dumps(row, ensure_ascii=False) + ''';
const scope='strategy_signals:202610', cold='strategy_events:202610';
const post=(data,action)=>doPost({postData:{contents:JSON.stringify({scope,data,action})}});
assert(post({model_trades:[original]}).success);
fail=true;assert(!post(null,'archive_model_events').success);
assert.equal(doGet({parameter:{scope}}).data.model_trades[0]['事件'].length,original['事件'].length);
assert(post(null,'archive_model_events').success);
const compact=doGet({parameter:{scope}}).data.model_trades[0];
const archive=doGet({parameter:{scope:cold}}).data;
assert(compact['事件'].length<10 && compact['資料缺口']);
assert.equal(compact['事件數'],original['事件'].length);
assert(compact['事件'].some(e=>e['事件ID']==='protect'));
assert.deepEqual(JSON.parse(archive.records[0].events_json),original['事件']);
assert.equal(modelDigest_(archive.records[0].events_json),compact['事件封存'].sha);
assert(!post({model_trades:[original]}).success);
assert(post({model_trades:[compact]}).success);
assert.equal(post(null,'archive_model_events').archived_trades,0);
assert(!doPost({postData:{contents:JSON.stringify({scope:cold,data:{records:[]}})}}).success);
console.log(JSON.stringify({compact,archive}));
'''
    result = subprocess.run([node], input=program, text=True, encoding='utf-8', capture_output=True)
    assert result.returncode == 0, result.stderr
    data = json.loads(result.stdout)
    tracker = ModelTracker(tmp_path/'archive.sqlite', start=False)
    with tracker.write_db() as db:
        tracker.store(db, row)
    late = {**row['事件'][-1], '事件ID':'late-stop', '狀態':'已平倉', '原因':'停損'}
    with tracker.write_db() as db:
        pending = {**row, '事件':row['事件']+[late]}
        tracker.store(db, pending)
    # Cloud archive reads happen outside the database write lock.
    def get(scope):
        with tracker.write_db():
            pass
        assert scope == 'strategy_events:202610'
        return data['archive']
    tracker.remote_get = get
    tracker.merge_synced_rows([data['compact']])
    compact = tracker.records()[0]
    assert len(compact['事件']) < 10 and compact['資料缺口']
    assert compact['事件數'] == len(row['事件']) + 1
    restored = tracker.details([row['交易ID']])[0]
    assert {e['事件ID']:e for e in restored['事件']} == {e['事件ID']:e for e in pending['事件']}
    assert tracker.db().execute('SELECT dirty FROM trades').fetchone()[0] == 1
    assert performance([row])[0] == performance([compact])[0]
    fresh = ModelTracker(tmp_path/'fresh.sqlite', start=False)
    fresh.merge_synced_rows([data['compact']])
    assert not fresh.events_ready([row['交易ID']], request=True)
    assert fresh.archive_requests == {'strategy_events:202610'}
    with pytest.raises(ValueError, match='載入'):
        fresh.details([row['交易ID']])
    fresh.remote_get = lambda scope: data['archive']
    fresh.load_event_archive('strategy_events:202610')
    assert fresh.events_ready([row['交易ID']]) and len(fresh.events([row['交易ID']], limit=500)) == 500
    assert len(fresh.details([row['交易ID']])[0]['事件']) == len(row['事件'])


def test_cross_device_reentry_identity_and_pending_ack_retry(tmp_path):
    devices = [ModelTracker(tmp_path/f'device{i}.sqlite', 'fake', start=False) for i in range(2)]
    evidence = [v12_quote(price, i)['成交證據'][0] for i, price in enumerate((100, 111, 99, 100))]
    for tracker in devices:
        tracker.process([signal(策略版本='V1.2-test')], {'股票|2330':v12_quote(100, 3, 成交證據=evidence)})
    rows = [tracker.records() for tracker in devices]
    assert {r['交易ID'] for r in rows[0]} == {r['交易ID'] for r in rows[1]} and len(rows[0]) == 2
    assert performance(rows[0])[0]['有效交易數'] == 0
    saved = [{**row, '同步確認':True} for row in merge_trades(*rows)]
    for tracker in devices:
        tracker.merge_synced_rows(saved, {r['交易ID']:r for r in tracker.records()})
        tracker.merge_synced_rows(saved)
        assert len(tracker.records()) == 2
        assert all(row['同步確認'] for row in tracker.records())
        assert performance(tracker.records())[0]['有效交易數'] == 1


@pytest.mark.parametrize('mode', ['當沖', '波段'])
def test_month_boundary_reentry_respects_plan_validity(tmp_path, mode):
    tracker = ModelTracker(tmp_path/(mode+'.sqlite'), start=False)
    start = '2026-09-30T23:59:58+08:00'
    import pandas as pd
    for i, price in enumerate((100,111,99,100)):
        time = (pd.Timestamp(start)+pd.Timedelta(seconds=i)).isoformat()
        day = time[:10]
        evidence = {'價格':price,'來源時間':time,'有效':True,'序號':i+1,'串流識別':'api:0'}
        quote = {**evidence,'交易日':day,'策略':mode,'策略版本':'V1.2-test', '成交證據':[evidence]}
        tracker.process([signal(交易日='2026-09-30',來源時間=start,策略=mode,策略版本='V1.2-test')] if i==0 else [], {'股票|2330':quote})
    rows = sorted(tracker.records(), key=lambda row:row['進場次數'])
    assert len(rows) == (2 if mode == '波段' else 1)
    if mode == '波段':
        assert rows[1]['前次執行月份'] == '202609' and rows[1]['交易日'] == '2026-10-01'


def test_expiry_and_explicit_new_permit_never_reuse_old_fill(tmp_path):
    tracker = ModelTracker(tmp_path/'permit.sqlite', start=False)
    tracker.process([signal(策略版本='V1.2-test')], {'股票|2330':v12_quote(100,0)})
    tracker.process([], {'股票|2330':v12_quote(111,1)})
    tracker.process([], {'股票|2330':v12_quote(100,2,進場許可識別='new',進場許可時間=quote(0,2)['來源時間'])})
    rows = sorted(tracker.records(), key=lambda row:row['進場次數'])
    assert len(rows) == 2 and rows[1]['成交證據識別']['來源時間'] == quote(0,2)['來源時間']
    expired = ModelTracker(tmp_path/'expiry.sqlite', start=False)
    expired.process([signal(策略版本='V1.2-test', 到期時間=quote(0,2)['來源時間'])], {'股票|2330':v12_quote(100,0)})
    expired.process([], {'股票|2330':v12_quote(111,1)})
    expired.process([], {'股票|2330':v12_quote(99,2)})
    expired.process([], {'股票|2330':v12_quote(100,3)})
    assert len(expired.records()) == 1
